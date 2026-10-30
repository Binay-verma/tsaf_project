"""Build a separate, audited explicit-feature results supplement."""
from pathlib import Path
import argparse
import json
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from build_multifestival_update import text, BG, MUTED

ROOT = Path(__file__).resolve().parents[1]
EVENT = ROOT/'results/festival_features_v1'
CONTROL = ROOT/'results/festival_feature_spillover_v2'
OUT = ROOT/'presentation/festival_feature_update_v1'
MODELS = ['ridge', 'lightgbm', 'lightgbm_capacity2', 'sarimax', 'chronos', 'timesfm']
NAMES = ['Rich Ridge', 'Stable LightGBM', 'LightGBM · leaf minimum 2', 'SARIMAX', 'Chronos-2', 'TimesFM-3']
ARMS = ['calendar', 'zero5', 'day5', 'window5', 'placebo_day5', 'placebo_window5']
ARM_NAMES = ['Ordinary\ncalendar', '+5 zeros', '+Festival\nday flags', '+Festival\nwindows', '+Placebo\nday flags', '+Placebo\nwindows']
FAMILIES = ['Holi', 'Dussehra', 'Ganesh_Chaturthi', 'Ugadi_Gudi_Padwa', 'Diwali']


def heatmap(table, title, output, column_labels, vmin=0, vmax=None, diverging=False):
    a = table.to_numpy(dtype=float)
    fig, ax = plt.subplots(figsize=(12, 4.5))
    im = ax.imshow(a, cmap='RdBu_r' if diverging else 'YlGnBu', vmin=vmin, vmax=vmax, aspect='auto')
    ax.set_yticks(range(len(table)), NAMES, fontsize=11)
    ax.set_xticks(range(len(table.columns)), column_labels, fontsize=11)
    ax.tick_params(length=0)
    for i in range(a.shape[0]):
        for j in range(a.shape[1]):
            if diverging:
                color = 'white' if abs(a[i, j]) > .65*vmax else 'black'
            else:
                color = 'white' if a[i, j] > .6*vmax else 'black'
            ax.text(j, i, f'{a[i, j]:+.3f}' if diverging else f'{a[i, j]:.3f}', ha='center', va='center', color=color, fontsize=12)
    ax.set_title(title, fontsize=14, pad=15)
    fig.colorbar(im, ax=ax, shrink=.8, label='Change in MASE (negative better)' if diverging else 'MASE (lower better)')
    fig.tight_layout(); fig.savefig(output, dpi=190); plt.close(fig)


def main(review_rebuild=False):
    if OUT.exists() and not review_rebuild:
        raise FileExistsError('Preserve earlier decks; version this supplement')
    for base in [EVENT, CONTROL]:
        for name in ['independent_audit.json', 'input_audit.json', 'fitting_audit.json']:
            assert json.loads((base/name).read_text())['status'] == 'pass'
    OUT.mkdir(parents=True, exist_ok=review_rebuild)
    m = pd.read_csv(EVENT/'primary_metrics_balanced.csv')
    c = pd.read_csv(CONTROL/'primary_metrics_balanced.csv')
    vmax = max(float(m.MASE.max()), float(c.MASE.max()))
    for h in [1, 7, 14]:
        e = m[m.horizon.eq(h)].pivot(index='model', columns='arm', values='MASE').reindex(index=MODELS, columns=ARMS)
        heatmap(e, f'Festival anchors · lead {h} · 2024–2025', OUT/f'event_lead_{h}.png', ARM_NAMES, vmax=vmax)
        e.rename(index=dict(zip(MODELS, NAMES))).to_csv(OUT/f'event_lead_{h}.csv')
        control = c[c.horizon.eq(h)].pivot(index='model', columns='arm', values='MASE').reindex(index=MODELS, columns=['calendar', 'day5', 'window5'])
        heatmap(control, f'Matched non-event targets · lead {h} · separate72-case cohort', OUT/f'non_event_lead_{h}.png', [ARM_NAMES[i] for i in [0, 2, 3]], vmax=vmax)
        control.rename(index=dict(zip(MODELS, NAMES))).to_csv(OUT/f'non_event_lead_{h}.csv')
    contrast = pd.read_csv(EVENT/'primary_paired_festival.csv')
    fig, axes = plt.subplots(1, 2, figsize=(13, 5.4), sharey=True)
    limit = float(contrast[contrast.horizon.eq(1) & contrast.arm.isin(['day5', 'window5'])].MASE_change.abs().max())
    limit = max(limit, 1e-3)
    for ax, arm, title in zip(axes, ['day5', 'window5'], ['Day flags minus calendar', 'Window flags minus calendar']):
        a = contrast[contrast.horizon.eq(1) & contrast.arm.eq(arm)].pivot(index='model', columns='family', values='MASE_change').reindex(index=MODELS, columns=FAMILIES)
        im = ax.imshow(a, cmap='RdBu_r', vmin=-limit, vmax=limit, aspect='auto')
        ax.set_yticks(range(6), NAMES, fontsize=10)
        ax.set_xticks(range(5), ['Holi', 'Dussehra', 'Ganesh', 'Ugadi/\nGudi', 'Diwali'], fontsize=10)
        ax.set_title(title, fontsize=12)
        for i in range(6):
            for j in range(5):
                v = a.iloc[i, j]
                ax.text(j, i, f'{v:+.2f}', ha='center', va='center', fontsize=10, color='white' if abs(v) > .65*limit else 'black')
    fig.suptitle('Next-day paired change in MASE · negative means lower error', fontsize=14)
    fig.subplots_adjust(left=.2, right=.85, top=.85, bottom=.18, wspace=.18)
    color_axis = fig.add_axes([.89, .18, .018, .67])
    fig.colorbar(im, cax=color_axis, label='Change in MASE')
    fig.savefig(OUT/'per_festival_day1_changes.png', dpi=190); plt.close(fig)
    d = pd.read_csv(EVENT/'predictions_secondary_classical.csv', parse_dates=['date'])
    examples = [('Holi', 'Delhi', 2024), ('Ganesh_Chaturthi', 'Maharashtra', 2025)]
    fig, axes = plt.subplots(1, 2, figsize=(13, 4.7))
    for ax, (family, state, year) in zip(axes, examples):
        g = d[d.model.eq('ridge') & d.family.eq(family) & d.state.eq(state) & d.year.eq(year) & d.horizon.eq(1)]
        base = g[g.arm.eq('calendar')].sort_values('date')
        ax.plot(base.date, base.actual, 'k-o', linewidth=2.6, label='Recorded actual')
        for arm, label in [('calendar', 'Ordinary calendar'), ('day5', '+ Day flags'), ('window5', '+ Windows'), ('offset35', '+ Relative-day channels')]:
            q = g[g.arm.eq(arm)].sort_values('date')
            ax.plot(q.date, q.prediction, '--', linewidth=1.6, label=label)
        ax.set_title(f'{state} · {family.replace("_", " ")} {year}', fontsize=12)
        ax.set_ylabel('Energy met/supplied (GWh)'); ax.tick_params(axis='x', rotation=35); ax.grid(alpha=.2)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc='lower center', ncol=3, fontsize=10)
    fig.suptitle('Ridge secondary windows · same two illustrations as prior supplement', fontsize=14)
    fig.tight_layout(rect=[0, .18, 1, .94]); fig.savefig(OUT/'ridge_window_examples.png', dpi=190); plt.close(fig)
    qa_event = json.loads((EVENT/'input_audit.json').read_text())
    qa_control = json.loads((CONTROL/'input_audit.json').read_text())
    crossings = sum(x['native_quantile_crossings'] for q in [qa_event, qa_control] for x in q['models'].values())
    slides = [
        dict(title='Do explicit festival features help?', body='Binay Verma · College results update · 1 November 2026\nAn exploratory added-feature experiment, with matched controls.\nDaily supplied electricity · Delhi, Karnataka and Maharashtra.',
            notes='This is distinct from the preceding diagnostic comparison, which measured errors during festivals without supplying all festival identities. The current experiment supplies family-specific day or window features. No new foundation-model architecture or clean final test is claimed.'),
        dict(title='What was checked between the steps?', body='Named training/context dates cover 2022–2025: 58 anchors.\nKarnataka Ganesh 2023: 18 Sep; Maharashtra: 19 Sep.\nSame-size zeros and shifted active placebo features.\nOriginal pre-year model settings retained; tree capacity checked separately.',
            notes='The Maharashtra2023official PDF and Karnataka2023English notice support state-specific training dates. The annual-as-issued ledger is not a full amendment or statewide shutdown certification. Rich Ridge and stable LightGBM reuse original ordinary-calendar hyperparameters for every arm. Leaf-minimum2is a fixed sensitivity with its own baseline, not an outcome-selected winning setting. Day flags may have only1–2historical positive examples.'),
        dict(title='Matched festival-anchor comparison', image=OUT/'event_lead_1.png',
            notes='Six model variants, six arms and72identical anchor cases at leads1/7/14. This slide shows lead1; all leads and individual families/years/states are saved. Lower MASE is better; it is not accuracy percent. Zero and placebo arms are included so a gain from adding channels alone is not assigned to event information. Two years and related states/days limit inference.'),
        dict(title='Effects differ across festival groups', image=OUT/'per_festival_day1_changes.png',
            notes='These are paired next-day MASE changes relative to the same model calendar baseline, with equal states within year then equal years. Negative values mean lower error. The two representations are separate: do not select whichever looks best for a clean test. Compare true features with zero/placebo controls in the full tables before attributing value to event information. Cohort state coverage differs by festival.'),
        dict(title='Check costs away from festivals', image=OUT/'non_event_lead_1.png',
            notes='Separate72date-selected non-event targets, matched by weekday and28–42day seasonal shifts. Every model uses identical cases. Control targets exclude +/-3days of annual/registered holiday anchors; source amendments remain incomplete. Family labels denote the matched event, not an observed holiday on the control date. Compare feature arms within this cohort; do not compare its raw error level with event errors as a causal effect. This is not an exhaustive all-day evaluation.'),
        dict(title='Inspect timing in the wider windows', image=OUT/'ridge_window_examples.png',
            notes='Only Ridge is plotted here on its secondary full-window cohort. The same Delhi Holi2024and Maharashtra Ganesh2025illustrations were used in the prior supplement; they were not newly selected for favorable intervention results. Every point is a rolling next-day forecast. Relative-day channels are a separate regression sensitivity, not a foundation-model experiment. Classical full-window scores must not be compared with foundation anchor-only scores.'),
        dict(title='What these results can establish', body='Whether these fixed festival encodings change matched forecast errors.\nWhether changes exceed the zero and active placebo controls.\nWhether ordinary-day errors also change under the same settings.\nModel, representation, sparse-history and year limits remain.',
            notes=f'All positive/negative outcomes, unsuccessful optimiser attempts and earlier results remain preserved. Across the new event/non-event outputs, {crossings}native adjacent-quantile crossings were retained without sorting or clipping. Point-error audit success does not establish calibrated uncertainty or independent human approval. No p-values or universal/causal winner claim is supported.'),
        dict(title='Sources and next scientific gates', body='Data: Robbie Andrew / GRID-INDIA / POSOCO; source clarification pending.\nCalendar evidence: annual state notices, with documented scope limits.\nModels: pinned Chronos, TimesFM and local classical/ML baselines.\nLater: source/human review and unchanged prospective outcome scoring.',
            notes='Cite the data curator, upstream organisation, model developers and close calendar/event-encoding prior work in the original presentation/references. This is not a new architectural method or first-ever holiday encoding. Source reuse response, amendment certification and independent human review remain open. The original31October–11Novemberprospective freeze is unchanged and cannot have a complete score on1November. These files are prepared locally; no new public upload is authorised by this supplement.')
    ]
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    prs.core_properties.author = 'Binay Verma'; prs.core_properties.title = 'Explicit festival features and matched controls'
    notes = ['# Explicit festival-feature experiment — Binay Verma', '', '1 November 2026 · Research in progress', '']
    for i, spec in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        text(slide, spec['title'], .55, .35, 12.2, .8, size=30, bold=True)
        if 'image' in spec:
            slide.shapes.add_picture(str(spec['image']), Inches(.45), Inches(1.4), width=Inches(12.4), height=Inches(5.2))
        else:
            text(slide, spec['body'], .7, 1.8, 11.9, 4.7, size=25)
        text(slide, f'Binay Verma | Exploratory added-feature study | 1 Nov 2026          {i}/{len(slides)}', .65, 7.05, 12, .3, size=11, color=MUTED)
        slide.notes_slide.notes_text_frame.text = spec['notes']
        notes += [f'## Slide {i}: {spec["title"]}', '', spec['notes'], '']
    path = OUT/'Festival_Features_Binay_Verma_2026-11-01.pptx'
    prs.save(path)
    (OUT/'SPEAKER_NOTES.md').write_text('\n'.join(notes)+'\n')
    (OUT/'README.md').write_text('# Explicit festival features: results supplement\n\nEight-slide local college supplement with matched event/ordinary-day comparisons, all-arm tables and recorded failure/uncertainty limits. Original presentations remain unchanged. See results/festival_features_v1 and results/festival_feature_spillover_v2 for full protocols, metrics and audits. No raw datasets or model weights embedded; no public upload performed.\n')
    print('Created8slides:', path)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--review-rebuild', action='store_true', help='Rebuild this current supplement during layout QA; scientific outputs and earlier decks stay unchanged')
    main(parser.parse_args().review_rebuild)
