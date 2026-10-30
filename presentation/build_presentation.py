"""Build an internal college deck from recorded aggregate research results."""
from pathlib import Path
import textwrap
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'presentation'
BG='101A2B';WHITE='F5F8FC';MUTED='B3C4D8';TEAL='36D3BD';AMBER='FFCA69'

def rgb(v):return RGBColor.from_string(v)

def label(slide,text,x,y,w,h,size=24,color=WHITE,bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.word_wrap=True
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.space_after=Pt(12)
    return box

def card(slide,title,body,x,y,w,h):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=rgb('1C2B42');shape.line.fill.background()
    label(slide,title,x+.22,y+.2,w-.44,.5,22,TEAL,True)
    label(slide,body,x+.22,y+.85,w-.44,h-.95,20)

r=ROOT/'results';ridge=pd.read_csv(r/'strong_refit_v1/metrics_macro.csv');tree=pd.read_csv(r/'stable_tree_refit_v2/metrics_macro.csv')
fig,axs=plt.subplots(1,2,figsize=(12,4),sharey=True)
for ax,model,data in [(axs[0],'Ridge',ridge[ridge.model.eq('ridge')]),(axs[1],'LightGBM (float32)',tree)]:
    d=data[data.year.eq(2026)&data.scope.eq('all')&data['mode'].eq('rich_calendar')]
    for schedule,g in d.groupby('schedule'):ax.plot(g.horizon,g.MASE,'o-',linewidth=2,label=schedule)
    ax.plot([1,7,14],[.467902,1.162937,1.361749],'o--',linewidth=2,label='Chronos history')
    ax.set_title(model);ax.set_xticks([1,7,14]);ax.set_xlabel('Individual lead day');ax.legend(fontsize=9)
axs[0].set_ylabel('Macro MASE (lower better)');fig.tight_layout();fig.savefig(OUT/'baseline_refits_for_slides.png',dpi=180);plt.close(fig)

slides=[
 dict(title='Do regional calendars help\nelectricity forecasting?',subtitle='Controlled evaluation on Indian state supplied-energy data',
      body='Binay Verma\nCollege research presentation · 1 November 2026\nResearch in progress',notes='Introduce the project as an empirical study. We use existing pretrained models and train local baselines; we do not claim a new foundation-model architecture. The presentation covers completed historical experiments, controls and data-quality findings. Prospective outcomes are still pending.'),
 dict(title='Why study this?',cards=[('Moving dates','Festival dates vary across years, and regional observances can differ.'),('Forecasting challenge','Pretrained models may use calendar inputs differently, even on the same cases.'),('Precise target','Daily energy met/supplied in GWh: what the system delivered, rather than latent unconstrained demand.')],
      notes='The question is whether regional date information adds predictive value after controlling the representation and feature count. Do not equate an office holiday with a proven statewide electricity response. Weather, economic activity and supply constraints can also affect the target; this is predictive analysis, not causal estimation.'),
 dict(title='Data and evaluation rules',bullets=['Three states: Karnataka, Delhi and Maharashtra.','Fixed source snapshot through 28 October 2026; sources remain separate.','730-day context; individual lead days 1, 7 and 14.','Chronological selection and purged labels; score recorded labels only.'],callout='Historical outcomes are exploratory; matched cases make comparisons meaningful.',
      notes='Explain origin as the first unseen date: lead h targets origin+h−1. Lead 7 is one target day, not a weekly average. We fill only historical inputs causally and do not score locally imputed labels. Retrospective revisions and unknown publication times mean date-order checks alone do not establish a fully operational real-time experiment. Cite Robbie Andrew and GRID-INDIA/POSOCO, with reuse clarification pending.'),
 dict(title='Models and comparison design',cards=[('Local models','Ridge and LightGBM\nRich causal history features\nAnnual versus monthly refits'),('Foundation models','Chronos-2 and TimesFM 3\nPinned pretrained weights\nHistorical context + known calendar'),('Fair comparisons','Same state/origin/lead cases\nPre-year parameter selection\nRegional transfer tested separately')],
      notes='Local hyperparameters use a pre-year temporal validation block and stay fixed for both deployment schedules during that year. The foundation models run pretrained inference; they are not trained from scratch here. Adaptation policies differ, so the benchmark is not an architecture-only experiment. Leave-region-out regression allows target-region historical context and normalization; it is not cold start.'),
 dict(title='Finding 1: refitting matters',image=OUT/'baseline_refits_for_slides.png',caption='Monthly updates improve several aggregates, but can worsen individual festival-year results.',source='Partial 2026, identical cases. Ridge version 1; tree float32 version 2; cached Chronos history.',
      notes='Use the chart to explain why a weak baseline is insufficient. With richer histories, monthly Ridge and LightGBM improve over their annual versions on several aggregate errors. Chronos remains lower on the displayed aggregate comparisons. Some event-year results become worse with monthly updates, so there is no universal winning refit schedule. Do not claim statistical superiority from these aggregate lines.'),
 dict(title='Controls before making a calendar claim',cards=[('Hold dimensions equal','Generic vs regional dates\nOne scalar window or seven offset channels'),('Add negative controls','One/seven zero channels\nActive weekday/season-matched placebo dates'),('Keep evidence fixed','Same 150 historical cases per arm\nTwo models, nine arms\nIndividual years and failures retained')],
      notes='The primary controlled event study covers 2024–2025. Two event-year clusters are not hundreds of independent experiments, even though we produce many forecasts. A single seeded placebo is not a randomization distribution. Scalar versus seven channels changes relative-day information as well as dimensionality; within-dimension generic/regional contrasts are more specific.'),
 dict(title='Finding 2: representation changes the answer',image=r/'mechanism_controls_v1/controlled_regional_effects.png',caption='Regional gains depend on the model, lead and representation.',source='Two historical event years; equal channel counts within each comparison; positive = lower MAE.',
      notes='For equal seven-channel inputs, regional mean state/year MAE improvements are about 1.4–1.6 percent for Chronos and 3.3–6.4 percent for TimesFM across leads. Scalar-window increments are smaller or mixed. These are exploratory effects, not universal or causal evidence. Adding seven zeros also changes forecasts by up to about 6–7 GWh, so interface/channel-count controls matter.'),
 dict(title='Quality checks found issues worth reporting',cards=[('Data revisions','Archived values differ from later snapshots by up to 61.5 GWh in the reviewed cases.'),('Reporting availability','Two tested issue-date snapshots lag by 2–3 calendar days; matched errors increase about 3.9–11.3%.'),('Numerical stability','Tiny rolling-arithmetic differences affect trees. A separate float32 contract reproduces the checked examples.')],callout='32 tests pass; separate-environment/software review is not external peer approval.',
      notes='These findings are part of the contribution, not defects to hide. Archive payload digests were verified. The selected festival labels agree across the checked later vintages, but some historical input values were revised. The availability experiment has only two dates. Original numerical failures and all old results remain preserved; the new stable tree run does not replace the frozen prospective forecasts.'),
 dict(title='Prospective validation: frozen before outcomes',timeline=True,caption='The full target period is not complete at this presentation.',
      notes='Forecasts were issued on 30 October using data through 28 October. Only 31 October–11 November targets qualify as future; 29/30 October are excluded. The GitHub workflow is registered and its pre-period dry run passed, but no actual future score exists yet. It checks time, source completeness and the original forecast Git blob before scoring. Do not report prospective accuracy on 1 November.'),
 dict(title='What the study supports today',bullets=['A reproducible empirical comparison with meaningful controls and retained negative results.','Model interfaces, adaptation policy and source vintages can change the interpretation.','Pooling regions can help; withholding the target region generally hurts the tested regression models.'],callout='A bounded research study—not a claim of a new architecture or universal festival gain.',
      notes='Present the completed work confidently but precisely. We found conditional calendar effects and failure modes, not a universally superior festival model. Three related states and two primary historical event years limit generalisation. Historic pretraining/proxy overlap remains unresolved. Data reuse clarification and independent human/calendar review are still pending.'),
 dict(title='What happens next?',cards=[('New outcomes','Score the immutable forecasts after 11 November, without retuning.'),('Human and source review','Permission email sent by the presenter; response pending. Independent reviewer still needed.'),('Submission decision','Review the evidence and limitations before deciding whether a journal submission is justified.')],
      notes='It is appropriate to discuss the project as work in progress in college. Use source acknowledgments and your own aggregate charts. Do not distribute raw datasets or model weights before clarifying terms. The author’s reply may take time. College feedback can inform improvements, but it is not automatically independent scientific sign-off or publication acceptance.'),
 dict(title='Sources and questions',bullets=['Robbie Andrew: Indian Energy and Emissions Data; upstream GRID-INDIA/POSOCO.','Hunt & Bloomfield: original corrected state dataset, Zenodo 14983362 (CC BY 4.0).','Chronos-2: arXiv 2510.15821; TimesFM: Google Research official repository/model.','Kreusel et al.: calendar encoding study, arXiv 2610.13676.'],callout='All displayed charts are generated from this project’s recorded aggregate analyses.',
      notes='Acknowledge the data curator, upstream organisation, model developers and close prior work. This is not a first-ever calendar-encoding claim. For follow-up, refer to the working manuscript, protocols and controlled-result reports. Avoid claiming a data licence or external review has already been granted. End with questions.')
]

prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
prs.core_properties.author='Binay Verma';prs.core_properties.title='Regional Calendars and Electricity Forecasting';prs.core_properties.subject='Internal college research presentation, 1 November 2026'
notes=['# Speaker notes — Binay Verma\n','Presentation date: **1 November 2026**. Suggested duration: 10–12 minutes.\n','Use as ongoing academic research. No prospective accuracy claim is available yet.\n']
for number,s in enumerate(slides,1):
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=rgb(BG)
 label(slide,s['title'],.6,.4,12.1,1.3 if number==1 else .8,34 if number==1 else 30,bold=True)
 line=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.6), Inches(1.38 if number>1 else 2.0), Inches(1.0), Inches(.055));line.fill.solid();line.fill.fore_color.rgb=rgb(TEAL);line.line.fill.background()
 if number==1:
  label(slide,s['subtitle'],.65,2.4,11.7,1.1,26,TEAL)
  label(slide,s['body'],.65,4.2,11.7,1.8,23)
 elif 'image' in s:
  slide.shapes.add_picture(str(s['image']),Inches(.6),Inches(1.65),width=Inches(12.1),height=Inches(4.04))
  label(slide,s['caption'],.65,5.95,12,.6,22,TEAL,True)
  label(slide,s['source'],.65,6.55,12,.3,11,MUTED)
 elif 'cards' in s:
  for i,(title,body) in enumerate(s['cards']):card(slide,title,body,.65+i*4.18,1.85,3.95,3.65)
  if s.get('callout'):label(slide,s['callout'],.7,5.95,11.8,.8,21,AMBER)
 elif s.get('timeline'):
  for i,(title,body) in enumerate([('30 October','Forecasts frozen in Git\nObserved context ends 28 Oct'),('31 Oct–11 Nov','Eligible target dates\nOutcomes still arriving'),('12 November onward','Date-gated collection/scoring\nNo prediction changes')]):card(slide,title,body,.65+i*4.18,1.85,3.95,3.5)
  label(slide,s['caption'],.7,5.9,11.9,.8,23,AMBER,True)
 else:
  y=1.8
  for bullet in s.get('bullets',[]):
   label(slide,'•  '+bullet,.7,y,11.8,.9,23);y+=1.03
  if s.get('callout'):label(slide,s['callout'],.7,6.05,11.8,.7,20,AMBER)
 label(slide,f'Binay Verma  |  Research in progress  |  1 Nov 2026                          {number:02d}/{len(slides):02d}',.65,7.05,12,.28,11,MUTED)
 if slide.notes_slide.notes_text_frame is not None:slide.notes_slide.notes_text_frame.text=s['notes']
 notes += [f'\n## Slide {number}: '+s['title'].replace('\n',' ')+'\n',s['notes']+'\n']
prs.save(OUT/'College_Presentation_Binay_Verma_2026-11-01.pptx')
(OUT/'SPEAKER_NOTES.md').write_text('\n'.join(notes))
print('Created',len(slides),'slides with speaker notes')
