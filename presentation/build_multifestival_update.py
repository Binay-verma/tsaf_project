"""Create a separate results supplement; preserve the original college deck."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT/'results/multifestival_models_v1'
OUT = ROOT/'presentation/multifestival_update_v1'
LABELS = {
    'chronos_history': 'Chronos · history', 'chronos_calendar': 'Chronos · calendar',
    'timesfm_history': 'TimesFM · history', 'timesfm_calendar': 'TimesFM · calendar',
    'sarimax_history_annual': 'SARIMAX · history · annual',
    'sarimax_calendar_annual': 'SARIMAX · calendar · annual',
    'ets_weekly_monthly': 'ETS · weekly · monthly',
    'ridge_basic_history': 'Ridge · basic history · annual',
    'ridge_basic_calendar': 'Ridge · basic calendar · annual',
    'ridge_rich_history_monthly': 'Ridge · rich history · monthly',
    'ridge_rich_calendar_monthly': 'Ridge · rich calendar · monthly',
    'lightgbm_stable_history_monthly': 'LightGBM · stable history · monthly',
    'lightgbm_stable_calendar_monthly': 'LightGBM · stable calendar · monthly',
    'last_value': 'Last value', 'weekly_repeat': 'Weekly repeat'}
FAMILIES = ['Holi', 'Dussehra', 'Ganesh_Chaturthi', 'Ugadi_Gudi_Padwa', 'Diwali']
FAMILY_LABELS = ['Holi\nDL, MH', 'Dussehra\nDL, KA, MH', 'Ganesh Chaturthi\nKA, MH', 'Ugadi/Gudi Padwa\nKA, MH', 'Diwali\nDL, KA, MH']
ORDER = list(LABELS)
BG, WHITE, TEAL, MUTED = '101A2B', 'F5F8FC', '36D3BD', 'B3C4D8'


def text(slide, value, x, y, w, h, size=22, color=WHITE, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    for i, line in enumerate(value.split('\n')):
        p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        p.text = line
        p.font.name = 'Arial'; p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(color); p.font.bold = bold
        p.space_after = Pt(12)


def main():
    if OUT.exists():
        raise FileExistsError('Preserve earlier deck; version the presentation update')
    for f in ['independent_audit.json', 'deployment_audit.json']:
        assert json.loads((RESULT/f).read_text())['status'] == 'pass'
    OUT.mkdir(parents=True)
    m = pd.read_csv(RESULT/'metrics_festival.csv')
    d = pd.read_csv(RESULT/'predictions.csv', parse_dates=['origin', 'date'])
    assert set(m.configuration) == set(ORDER)
    for h in [1, 7, 14]:
        table = m[m.horizon.eq(h)].pivot(index='configuration', columns='family', values='MASE').reindex(index=ORDER, columns=FAMILIES)
        fig, ax = plt.subplots(figsize=(13, 6.3))
        # The color scale is fixed across horizons so colors are comparable.
        im = ax.imshow(table.to_numpy(), cmap='YlGnBu', vmin=0, vmax=2.5, aspect='auto')
        ax.set_yticks(range(len(ORDER)), [LABELS[x] for x in ORDER], fontsize=10)
        ax.set_xticks(range(5), FAMILY_LABELS, fontsize=10)
        ax.tick_params(length=0)
        for i in range(len(ORDER)):
            for j in range(5):
                v = table.iloc[i, j]
                ax.text(j, i, f'{v:.3f}', ha='center', va='center', fontsize=10, color='white' if v > 1.3 else 'black')
        ax.set_title(f'Lead {h}: MASE, lower is better | 2024–2025 matched cases', fontsize=14, pad=12)
        fig.colorbar(im, ax=ax, shrink=.8, label='MASE (scale fixed across leads)')
        fig.tight_layout()
        fig.savefig(OUT/f'lead_{h}_all_configurations.png', dpi=180)
        plt.close(fig)
        table.rename(index=LABELS).to_csv(OUT/f'lead_{h}_table.csv')
    # Prespecified illustrations: not the largest errors or best model examples.
    examples = [('Holi', 'Delhi', 2024), ('Ganesh_Chaturthi', 'Maharashtra', 2025)]
    show = ['chronos_history', 'timesfm_calendar', 'sarimax_calendar_annual', 'ridge_rich_calendar_monthly', 'ets_weekly_monthly']
    fig, axes = plt.subplots(1, 2, figsize=(13, 4.8))
    for ax, (family, state, year) in zip(axes, examples):
        g = d[d.family.eq(family) & d.state.eq(state) & d.year.eq(year) & d.horizon.eq(1)]
        a = g[g.configuration.eq('chronos_history')].sort_values('date')
        ax.plot(a.date, a.actual, 'k-o', label='Recorded actual', linewidth=2.5)
        for config in show:
            q = g[g.configuration.eq(config)].sort_values('date')
            ax.plot(q.date, q.prediction, '--', label=LABELS[config], linewidth=1.4)
        ax.set_title(f'{state} · {family.replace("_", " ")} {year}')
        ax.set_ylabel('Daily energy met/supplied (GWh)')
        ax.tick_params(axis='x', rotation=35)
        ax.grid(alpha=.2)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc='lower center', ncol=3, fontsize=9)
    fig.suptitle('Rolling next-day forecasts: two date-selected illustrations', fontsize=14)
    fig.tight_layout(rect=[0, .17, 1, .94])
    fig.savefig(OUT/'rolling_next_day_examples.png', dpi=180)
    plt.close(fig)
    findings = []
    for family in FAMILIES:
        parts = []
        for h in [1, 7, 14]:
            g = m[m.family.eq(family) & m.horizon.eq(h)]
            r = g.loc[g.MASE.idxmin()]
            parts.append(f'Lead {h}: {LABELS[r.configuration]} ({r.MASE:.3f})')
        findings.append(f"{family.replace('_', ' ')}: "+'; '.join(parts))
    (OUT/'OBSERVED_LOWEST_ERRORS.md').write_text('# Descriptive lowest errors among15configurations\n\nThese are observed historical minima, not selected future models or statistical superiority claims. Consult all-model tables.\n\n'+'\n\n'.join(findings)+'\n')
    slides = [
        dict(title='Do rankings hold beyond Diwali?', body='Binay Verma · Results supplement · 1 November 2026\nFive festival groups · 2024–2025 · Research in progress',
            notes='This supplement extends the earlier Diwali study. It evaluates forecast errors around five documented festival groups. It does not establish a new architecture, causal effect or benefit from adding new festival identities.'),
        dict(title='Expanded comparison, fixed evaluation cases', body='528 identical cases per configuration; 15 configurations.\nHoli: Delhi/Maharashtra. Ganesh and Ugadi/Gudi: Karnataka/Maharashtra.\nDussehra and Diwali: all three states. Fixed windows and recorded labels.\nNaive, Ridge, LightGBM, SARIMAX, ETS, Chronos and TimesFM families.',
            notes='There are24festival/state/year groups, not528independent replications. Diwali preserves the earlier broad union; other groups use fixed plus/minus3day anchor windows. Fifteen configurations have different adaptation and feature budgets, so this is a deployment-policy comparison. SARIMAX and ETS use fixed specifications; learned Ridge/tree settings retain their original pre-year selection. Date evidence is annual-as-issued.'),
        *[dict(title=f'Lead {h}: all configurations', image=OUT/f'lead_{h}_all_configurations.png',
            notes='Lower MASE is better; these are not percentages. Each festival uses equal states within year and equal years. Different families have different documented state coverage. Every configuration is shown, including poor results. Colors use the same scale across leads. A lead is an individual future day, not an average over that many days. Do not infer statistical superiority from the lowest cell.') for h in [1, 7, 14]],
        dict(title='Inspect forecasts, not just rankings', image=OUT/'rolling_next_day_examples.png',
            notes='These illustrations were selected by dates and states before plotting, not because they are the most favorable errors: Delhi Holi2024 and Maharashtra Ganesh2025. Each point is a separately issued next-day prediction using earlier observations; it is not a single seven-day forecast path. The five plotted configurations are fixed for readability; all fifteen are retained in the tables.'),
        dict(title='What the extension supports', body='Model rankings depend on festival, state, year and horizon.\nClassical models and simple baselines remain relevant comparators.\nOrdinary-calendar inputs do not encode every festival identity.\nLow historical error does not establish a universal winner.',
            notes='This is an exploratory diagnostic result. We have extended case coverage, not proved the incremental value of adding every festival feature. Calendar arms here contain weekdays and annual Fourier harmonics. Foundation-model historical/proxy overlap remains unresolved; matching cases controls a specific comparison, not every source of bias.'),
        dict(title='Next experiment and remaining limits', body='Next: test added festival features using a separate training-date ledger.\nTwo historical years, correlated states and overlapping holidays limit inference.\nSource permission, amendment review and independent human review remain pending.\nFrozen future outcomes: score only after 11 November; no accuracy claim yet.',
            notes='Dussehra2025 coincides with Gandhi Jayanti and Spring2025 is adjacent to Eid. Annual office schedules do not establish actual statewide shutdown or all later amendments. Cite Robbie Andrew/GRID-INDIA/POSOCO, the earlier Hunt/Bloomfield source, model developers and close calendar-encoding prior work. Sources and full scores are in multifestival_models_v1 and MULTIFESTIVAL_REPRODUCTION.md. The original college deck and frozen predictions remain unchanged. This supplement is local pending sharing clarification.')
    ]
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    prs.core_properties.author = 'Binay Verma'
    prs.core_properties.title = 'Multi-festival forecasting: matched historical results'
    notes = ['# Multi-festival results supplement — Binay Verma', '', '1 November 2026 · Research in progress', '']
    for number, spec in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(BG)
        text(slide, spec['title'], .55, .35, 12.2, .8, size=30, bold=True)
        if 'image' in spec:
            height = 5.65 if number in [3, 4, 5] else 4.9
            slide.shapes.add_picture(str(spec['image']), Inches(.45), Inches(1.18), width=Inches(12.4), height=Inches(height))
        else:
            text(slide, spec['body'], .7, 1.85, 11.9, 4.7, size=25)
        text(slide, f'Binay Verma | Exploratory historical results | 1 Nov 2026                 {number}/{len(slides)}', .65, 7.05, 12, .3, size=11, color=MUTED)
        slide.notes_slide.notes_text_frame.text = spec['notes']
        notes += [f'## Slide {number}: {spec["title"]}', '', spec['notes'], '']
    path = OUT/'College_Results_Update_Binay_Verma_2026-11-01.pptx'
    prs.save(path)
    (OUT/'SPEAKER_NOTES.md').write_text('\n'.join(notes)+'\n')
    (OUT/'README.md').write_text('# Multi-festival college results update\n\nEight-slide supplement to the original12-slide deck. Uses only project-generated aggregate tables/charts. No raw datasets, model weights or source-document scans embedded. Full source, protocols and audits remain in results/multifestival_models_v1 and MULTIFESTIVAL_REPRODUCTION.md. Preserve the original deck; these files are local pending sharing clarification.\n')
    print('Built', len(slides), 'slides:', path)


if __name__ == '__main__':
    main()
