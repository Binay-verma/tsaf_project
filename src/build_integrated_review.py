"""Integrate complete, audited experiments into a NEW review draft.

This is reporting, not another model-selection step. All methods, leads and
controls remain in supplementary tables. Never substitutes local checkpoints
for missing foundation outputs or modifies an older manuscript.
"""
from datetime import datetime, timezone
import json
from pathlib import Path
import subprocess

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document

from build_journal_draft import format_docx, NAMES
from festival_feature_design import ROOT, dump, sha
from report_multifestival import markdown

OUT = ROOT / 'manuscript/integrated_review_v3'
RECENTER = ROOT / 'results/calendar_recenter_recovery_v1/analysis'
MODELS = ['ridge', 'sarimax', 'lightgbm', 'lightgbm_capacity2', 'chronos', 'timesfm']
ARMS = ['calendar', 'day5', 'pooled_day5', 'window5', 'pooled_window5']
ARM_LABELS = {'calendar': 'Calendar', 'day5': 'Typed day',
              'pooled_day5': 'Pooled day', 'window5': 'Typed window',
              'pooled_window5': 'Pooled window'}
GATES = [
    'results/feature_breadth_v1/independent_audit.json',
    'results/calendar_encoding_v1/independent_audit.json',
    'results/neso_replication_v1/analysis/independent_audit.json',
    'results/calendar_semantics_v2/analysis/independent_audit.json',
    'results/calendar_recenter_recovery_v1/analysis/independent_audit.json',
    'results/calendar_recenter_recovery_v1/analysis/contrasts/audit.json',
    'results/mta_holiday_v1/analysis/independent_audit.json',
    'results/event_comparators_v1/independent_audit.json',
    'results/pooling_simulation_v1/independent_audit.json',
    'review/clean_extended_v1/comparison.json',
]


def formatted(value):
    if isinstance(value, (float, np.floating)):
        return 'undefined' if pd.isna(value) else f'{value:.5f}'
    return str(value)


def formatted_frame(frame):
    return pd.DataFrame([[formatted(v) for v in row]
                         for row in frame.itertuples(index=False, name=None)],
                        columns=frame.columns)


def main():
    if OUT.exists():
        raise FileExistsError('Preserve prior integrated draft; version recovery')
    for name in GATES:
        assert json.loads((ROOT/name).read_text())['status'] == 'pass', name

    clean_mta_path = ROOT/'review/clean_mta_v1/comparison.json'
    clean_mta = json.loads(clean_mta_path.read_text())
    assert clean_mta['status'] == 'pass' and clean_mta['rows'] == 4032
    assert all(item['prediction_csv_byte_identical'] for item in clean_mta['workers'].values())
    support_audit_path = ROOT/'results/event_support_v1/independent_audit.json'
    support_audit = json.loads(support_audit_path.read_text())
    assert support_audit['status'] == 'pass' and support_audit['paired_cases'] == 6336

    # Load everything BEFORE creating outputs. An incomplete study is a hard stop.
    sources = {}

    def read(path):
        path = Path(path)
        frame = pd.read_csv(path)
        sources[str(path.relative_to(ROOT))] = sha(path)
        return frame

    india = {cohort: read(RECENTER/'breadth'/cohort/'metrics_balanced.csv')
             for cohort in ['event_window', 'ordinary']}
    encoding = {cohort: read(RECENTER/'encoding'/cohort/'metrics_balanced.csv')
                for cohort in india}
    native = read(RECENTER/'breadth/event_window/native_uncertainty_balanced.csv')
    paired = {cohort: read(RECENTER/'contrasts/breadth'/cohort/'balanced.csv')
              for cohort in india}
    mta_root = ROOT/'results/mta_holiday_v1/analysis'
    mta = {cohort: read(mta_root/cohort/'metrics_balanced.csv') for cohort in india}
    neso = read(ROOT/'results/neso_replication_v1/analysis/metrics.csv')
    comparators = read(ROOT/'results/event_comparators_v1/metrics_balanced.csv')
    simulation = read(ROOT/'results/pooling_simulation_v1/paired_seed_contrasts.csv')
    changes = read(RECENTER/'breadth/event_window/changes_balanced.csv')
    composition = read(RECENTER/'breadth/event_window/composition_changes.csv')
    support = read(ROOT/'results/event_support_v1/associations.csv')
    for name in ['moderators.csv', 'paired_cases.csv', 'cells.csv']:
        read(ROOT/'results/event_support_v1'/name)
    for collection in [india, encoding]:
        for cohort, frame in collection.items():
            selected = frame[frame.version.eq('document_aligned')]
            expected = len(MODELS)*8*3 if collection is india else 2*11*3
            assert len(selected) == expected, (cohort, expected, len(selected))
            assert not selected.duplicated(['model', 'arm', 'horizon']).any()
    assert set(mta['event_window'].model) == set(MODELS+['last', 'seasonal'])

    OUT.mkdir()
    (OUT/'figures').mkdir()
    (OUT/'supplement').mkdir()
    tables = []

    def add_table(name, title, frame):
        assert not frame.empty
        frame.to_csv(OUT/'supplement'/(name+'.csv'), index=False)
        tables.append(dict(name=name, title=title, frame=frame))
        return title+'\n\n'+markdown(formatted_frame(frame))+'\n'

    def main_scores(frame, column='horizon', lead=1):
        if 'version' in frame:
            frame = frame[frame.version.eq('document_aligned')]
        part = frame[frame[column].eq(lead)&frame.arm.isin(ARMS)]
        matrix = part.pivot(index='model', columns='arm', values='MASE')
        matrix = matrix.reindex(index=MODELS, columns=ARMS)
        assert np.isfinite(matrix.to_numpy()).all()
        matrix.index = [NAMES[m] for m in MODELS]
        return matrix.rename(columns=ARM_LABELS).rename_axis('Model').reset_index()

    main_tables = {}
    main_tables['india_event'] = add_table('india_event',
        'Table 1. Document-aligned, recentered Indian festival windows: lead-1 MASE.',
        main_scores(india['event_window']))
    main_tables['india_ordinary'] = add_table('india_ordinary',
        'Table 2. The matched Indian seasonal ordinary sample: lead-1 MASE.',
        main_scores(india['ordinary']))
    pure = encoding['event_window']
    pure = pure[pure.version.eq('document_aligned')&pure.horizon.eq(1)]
    order = ['history', 'zero1', 'ordinal', 'onehot', 'cyclic', 'te_weekday',
             'te_sunday', 'te_holiday', 'te_joint', 'te_separate', 'calendar']
    # Use the registration's actual names; fail rather than infer a missing row.
    assert set(pure.arm) == set(order), sorted(pure.arm.unique())
    pure = pure.pivot(index='arm', columns='model', values='MASE').reindex(order)
    pure = pure.reindex(columns=['chronos', 'timesfm']).rename(columns=NAMES)
    main_tables['pure'] = add_table('pure',
        'Table 3. Matched Indian festival windows: every registered pure encoding and calendar reference, lead-1 MASE.',
        pure.rename_axis('Encoding').reset_index())
    main_tables['mta_event'] = add_table('mta_event',
        'Table 4. NYC event windows: decision lead 1, actual model step 8, MASE.',
        main_scores(mta['event_window'], 'decision_lead'))
    n = neso[neso.decision_lead.eq(1)&neso.arm.eq('calendar')]
    n = n.pivot(index='model', columns='scenario', values='MASE').reindex(MODELS)
    assert np.isfinite(n.to_numpy()).all()
    n.index = [NAMES[m] for m in MODELS]
    main_tables['neso'] = add_table('neso',
        'Table 5. NESO calendar models at decision lead 1: ideal model step 1 versus lag-21 model step 22, MASE.',
        n.rename(columns={'ideal': 'Ideal', 'lag21': 'Lag 21'}).rename_axis('Model').reset_index())
    q = native[native.version.eq('document_aligned')&native.horizon.eq(1)&native.arm.isin(ARMS)].copy()
    assert len(q) == 10
    q['model'] = q.model.map(NAMES)
    q['arm'] = q.arm.map(ARM_LABELS)
    q = q[['model', 'arm', 'coverage80', 'scaled_width80', 'scaled_WIS80_median']]
    q.columns = ['Model', 'Arm', 'Coverage 80%', 'Scaled width', 'Scaled WIS80']
    main_tables['native'] = add_table('native',
        'Table 6. Unmodified native Indian festival-window intervals at lead 1. Undefined interval cells retain invalid endpoints.', q)
    appendix = add_table('mta_ordinary',
        'Table A1. NYC ordinary sample: decision lead 1, actual model step 8, MASE.',
        main_scores(mta['ordinary'], 'decision_lead'))
    c = comparators[comparators.cohort.eq('event_window')&comparators.method.isin(
        ['response_none', 'response_typed', 'response_pooled', 'response_partial'])]
    c = c.pivot(index='method', columns='horizon', values='MASE')
    c.columns = [f'Lead {h}' for h in c.columns]
    appendix += '\n'+add_table('response_comparators',
        'Table A2. Fixed response-profile comparators on the original Indian case ledger. This is a separate, unchanged sensitivity.',
        c.rename_axis('Method').reset_index())
    rows = []
    for model in MODELS:
        records = []
        for metric in ['own_support_ratio', 'heterogeneity_scaled']:
            item = support[support.model.eq(model)&support.representation.eq('day')&
                support.horizon.eq(1)&support.scope.eq('all')&support.moderator.eq(metric)]
            assert len(item) == 1
            records += [item.spearman.iloc[0], str(int(item.defined_cells.iloc[0]))]
        rows.append([NAMES[model]]+records)
    appendix += '\n'+add_table('support_diagnostics',
        'Table A3. Prespecified day-representation, lead-1 Spearman diagnostics. Positive forecast contrasts favour pooling; n counts available family/state/year cells, not independent replications. Undefined means insufficient variation or support.',
        pd.DataFrame(rows, columns=['Model', 'Support rho', 'Support n', 'Response rho', 'Response n']))

    # Export full audited metric/contrast tables, never just a favourable subset.
    source_index = {}
    for index, (relative, expected) in enumerate(sources.items()):
        assert sha(ROOT/relative) == expected
        target = OUT/'supplement'/f'source_{index:02d}_{Path(relative).name}'
        target.write_bytes((ROOT/relative).read_bytes())
        source_index[str(target.relative_to(OUT))] = dict(original=relative, sha256=expected)
    dump(OUT/'supplement/source_index.json', source_index)
    future_text = 'A separately registered Chronos history/calendar competitor follows the existing provider lane. Its new predictions/public issuance are not assumed complete here; it does not test festival features.'
    future_inputs = {}
    future_root = ROOT/'results/prospective_chronos_20261031_v1'
    if all((future_root/name).exists() for name in ['freeze.json', 'independent_audit.json', 'public_receipt.json']):
        from score_chronos_competitor import first_commit_guard
        first_commit_guard()  # no labels/network; checks actual completion/public-receipt chronology
        frozen = json.loads((future_root/'freeze.json').read_text())
        audited = json.loads((future_root/'independent_audit.json').read_text())
        assert audited['status'] == 'pass' and frozen['rows'] == 84 and frozen['eligible_rows'] == 66
        future_text = f'A new, separately timed Chronos history/calendar competitor completed at {frozen["issued_utc"]}, with 84 stored forecasts and 66 eligible full 1–11 November cases. Its observed pre-outcome public receipt and separate input/India-day audit are retained. It uses the same stale source/context as the older episode but a later actual issuance; the paired comparison will reuse the original first actual vintage. It fills the missing Chronos competitor, not the prospective festival-feature gap or an independent episode.'
        future_inputs = {str((future_root/name).relative_to(ROOT)): sha(future_root/name)
                         for name in ['protocol.json', 'freeze.json', 'independent_audit.json', 'public_receipt.json']}

    def difference(frame, model, arm, reference, lead, lead_column):
        if 'version' in frame:
            frame = frame[frame.version.eq('document_aligned')]
        subset = frame[frame.model.eq(model)&frame[lead_column].eq(lead)]
        a = subset[subset.arm.eq(arm)].MASE
        b = subset[subset.arm.eq(reference)].MASE
        assert len(a) == len(b) == 1
        return float(a.iloc[0]-b.iloc[0])

    def heatmaps(collection, lead_column, output, titles):
        arrays = {}
        for cohort, frame in collection.items():
            values = [[difference(frame, model, arm, reference, lead, lead_column)
                       for arm, reference in [('day5', 'pooled_day5'), ('window5', 'pooled_window5')]
                       for lead in [1, 7, 14]] for model in MODELS]
            arrays[cohort] = np.asarray(values)
        limit = max(float(abs(values).max()) for values in arrays.values())
        limit = max(limit, 1e-10)
        fig, axes = plt.subplots(2, 1, figsize=(8.6, 7.1), layout='constrained')
        for ax, (cohort, values) in zip(axes, arrays.items()):
            im = ax.imshow(values, cmap='RdBu_r', vmin=-limit, vmax=limit, aspect='auto')
            ax.set_yticks(range(6), [NAMES[m] for m in MODELS], fontsize=9)
            ax.set_xticks(range(6), ['Day 1', 'Day 7', 'Day 14', 'Window 1', 'Window 7', 'Window 14'], fontsize=9)
            ax.set_title(titles[cohort], fontsize=11)
            ax.axvline(2.5, color='white', linewidth=3)
            for i in range(6):
                for j in range(6):
                    ax.text(j, i, f'{values[i,j]:+.3f}', ha='center', va='center', fontsize=9,
                            color='white' if abs(values[i,j]) > .65*limit else 'black')
        fig.colorbar(im, ax=axes, label='Typed minus pooled MASE; negative favours typed')
        for extension in ['png', 'pdf']:
            fig.savefig(OUT/'figures'/(output+'.'+extension), dpi=180, bbox_inches='tight')
        plt.close(fig)
        return arrays

    arrays = heatmaps(india, 'horizon', 'india_contrasts',
        {'event_window': 'Indian complete festival windows', 'ordinary': 'Indian seasonal ordinary sample'})
    mta_arrays = heatmaps(mta, 'decision_lead', 'mta_contrasts',
        {'event_window': 'NYC complete event windows', 'ordinary': 'NYC weekday-balanced ordinary sample'})
    lo, hi = float(arrays['event_window'].min()), float(arrays['event_window'].max())
    mlo, mhi = float(mta_arrays['event_window'].min()), float(mta_arrays['event_window'].max())
    h = changes[changes.horizon.eq(1)].MASE_change
    dc = composition[composition.horizon.eq(1)].MASE_composition_change
    claims = {'india_min': lo, 'india_max': hi, 'mta_min': mlo, 'mta_max': mhi,
              'aligned_feature_change_min': float(h.min()), 'aligned_feature_change_max': float(h.max()),
              'composition_change_min': float(dc.min()), 'composition_change_max': float(dc.max())}

    # The synthetic figure is explicitly a generator illustration, not real-event uncertainty.
    s = simulation[simulation.arm.eq('typed_OLS')&simulation.reference.eq('pooled_OLS')]
    fig, axes = plt.subplots(1, 2, figsize=(9, 3.5), layout='constrained', sharey=True)
    for ax, heterogeneity in zip(axes, [0., 3.]):
        part = s[s.heterogeneity.eq(heterogeneity)].sort_values(['frequency', 'date_error', 'noise'])
        ax.errorbar(np.arange(len(part)), part.event_MAE_difference,
                    yerr=1.96*part.MC_standard_error, fmt='o', capsize=3)
        ax.axhline(0, color='black', linewidth=1)
        ax.set_title('Homogeneous' if heterogeneity == 0 else 'Heterogeneous')
        ax.set_xticks(np.arange(len(part)),
                      [f'{int(r.frequency)}/{int(r.date_error)}/{int(r.noise)}' for r in part.itertuples()],
                      rotation=45, fontsize=8)
        ax.set_xlabel('Frequency / date-error bound / noise SD')
    axes[0].set_ylabel('Typed minus pooled event MAE\n(95% Monte Carlo intervals)')
    for extension in ['png', 'pdf']:
        fig.savefig(OUT/'figures'/('simulation.'+extension), dpi=180, bbox_inches='tight')
    plt.close(fig)

    text = f'''# When does pooling sparse event information help forecasting? Matched representations, calendar corrections and cross-source evidence

**Integrated research draft for professor review. Not submitted or submission-ready.**

Binay Verma. College/department: __________. Corresponding-author email: __________.

## Abstract

Sparse events create a familiar trade-off between separate response identities and statistical sharing. Covariate-informed foundation models add sensitivity to channel allocation and encoding. We study this conditional forecasting problem using two pinned foundation models and four local variants, retaining calendar, zero-channel, typed, pooled and weekday-placebo controls. The main analysis uses 528 complete festival-window cases and 216 seasonal ordinary cases from three Indian states and two historical years. A source-driven correction changes one named festival date, restores its complete window and weekday-matched placebo, and separates feature changes from case-composition changes. Closest-method daily calendar encodings, established response comparators and a controlled simulation supplement this analysis. A pre-acquisition registered NYC replication adds two later holiday occurrences from a separate ridership lineage; an England/Wales electricity study examines a different source and an explicit observation-delay sensitivity. Typed-minus-pooled event MASE spans {lo:.5f} to {hi:.5f} in the aligned Indian panel and {mlo:.5f} to {mhi:.5f} in NYC. These are complete registered contrast ranges, not selected improvements or sampling intervals. Source revisions, sparse event years, unmatched adaptation policies, model exposure, unresolved Indian reuse terms and absent independent human review limit generalisation. The evidence qualifies when representations help; it does not introduce a new architecture or establish causal event effects.

**Keywords:** event representation; pooling; foundation models; calendar provenance; electricity; temporal evaluation.

## 1. Question, contribution and closest work

When does sharing an event response across sparse named families help, and does that answer survive the model interface, ordinary-day costs and a new source? Pooling, event regressors, partial pooling and calendar encodings are established ideas. A useful contribution must show conditions and failure modes beyond adding models to a benchmark.

Kreusel et al. [1] directly examine calendar encodings in Chronos-2 and TimesFM-3. Alipour et al. [2] learn impact-driven event representations from residual responses. Indian moving-holiday regression [3], covariate-uncertainty evaluation [4] and sample/proxy leakage analysis [5] are close antecedents. We compare simple representations under logged information and channel contracts rather than claim to invent them. The proposed contribution is a conditional empirical qualification with complete controls, source-driven corrections and separate-source evidence. Whether it is sufficiently distinct and general for International Journal of Forecasting remains an external scientific judgment.

The primary manuscript tables now use the document-aligned, recentered panel. Original forecasts, the original-ledger manuscript, the flag-only correction, failed attempts and negative findings remain available. A calendar repair is justified by its source, regardless of its forecast score. It is not permission to select the better historical version.

## 2. Data, dates and forecast information

### 2.1 Indian supplied electricity and named calendars

The target is daily energy met/supplied in GWh for Delhi, Karnataka and Maharashtra. It is not latent unconstrained demand, household consumption or hourly peak load. The later GRID-INDIA/POSOCO compilation is kept separate from an earlier corrected source. Outcomes through 28 October 2026 have been inspected; new historical method choices are exploratory. Historical revisions, exact release availability and correlated pretraining proxies remain unresolved. Weight dates alone do not establish contamination-free evaluation [5–7].

Five families are Holi, Dussehra, Ganesh Chaturthi, Ugadi/Gudi Padwa and Diwali, with declared state scopes and subsidiary anchors. Twelve retained annual documents support an internal all-58 named-date review: 54 declared-family matches, two explicit shared-national-main exceptions, one Karnataka 2025 mismatch and one Maharashtra 2022 Gudi Padwa evidence gap. The last means the retained one-page schedule does not corroborate the name/date; it is not proof that 2 April is wrong. Visible notification dates are not certified first-publication histories, and amendment chains and human sign-off remain incomplete. Office closure does not establish statewide electricity response.

The Karnataka notification labels Vijayadasami on 1 October 2025 and Gandhi Jayanthi on 2 October. The original ledger used 2 October for Dussehra. A separately registered flag-only correction holds cases and old placebos fixed. A further registered recentering replaces only three old 5 October target/lead cases with 28 September, yielding the complete 28 September–4 October window, and moves the corresponding +35-day placebo from 6 to 5 November. No error selects these dates. Both original-feature and aligned-feature forecasts are compared on the identical corrected cohort. Case-composition changes under original features are reported separately.

The resulting 528 event-window cases are 176 state/target-date combinations at leads 1, 7 and 14 across 24 family/state/year groups in 2024–2025. Nearby days, states and leads are correlated, not independent event replications. The 216 ordinary cases are a bounded seasonal sample: the first eligible rotating weekday per month/state/year, excluding registered windows. This is not exhaustive daily deployment evaluation.

### 2.2 NYC: a separately registered holiday replication

The canonical MTA daily ridership/traffic dataset, identifier sayj-mze2, provides Subway and Bus reported estimated counts. A fixed query returned 2,780 records and 1,390 complete daily labels per mode, 1 January 2023–21 October 2026. Counts are converted to million estimated rides, not GWh or unique travellers. The two modes are related series in one city and one source lineage; no deprecated-view splice or Staten Island addition is made. Actual API mode type is TEXT, although the attached dictionary says NUMERIC; the discrepancy is retained.

The protocol was publicly committed before these new evaluation records were acquired. Five calendar families retain the Indian study's five-channel budget. The 84 event cases cover complete ±3-day windows around Yom Kippur on 21 September and Columbus Day on 12 October 2026, two modes and three decision leads. The 42 ordinary cases balance weekdays and exclude all registered September–October NYC DOT observance windows. The calendar PDFs have matching archive digests in pre-year captures. They document parking observances, not MTA service closures or complete amendments. These target dates follow both actual model-weight uploads, but correlated proxy exposure is not ruled out.

A fixed seven-day observation buffer maps decision leads 1/7/14 to actual model steps 8/14/21. The intervening labels are unobserved, unscored bridge dates. This sensitivity accommodates uncertain update timing without pretending to reconstruct first releases. The current source is revised; the study is retrospective despite registration before target acquisition. Two occurrences cannot establish broad new-geography or annual generalisation.

### 2.3 England/Wales: source and delay, not holiday confirmation

NESO transmission demand supplies an independent electricity lineage. Daily GWh integrate complete actual half-hourly MW averages, including 46/48/50-period clock changes; forecast flags and incomplete days are excluded. Sixteen ordinary targets, 15–30 September 2026, are compared under ideal availability and 21 additional days of observation delay. Decision leads 1/7/14 become actual steps 22/28/35 in the delayed case. The archived estimates can be revised and published in arrears. This is a fixed lag sensitivity, not a first-vintage reconstruction, and it contains no target holiday. It cannot replace the NYC holiday replication.

### 2.4 Scoring and dependence

Each origin is the first unseen model date. Target date equals origin plus model step minus one. Scales are fixed pre-year 730-calendar-day mean absolute weekly differences, using positive recorded labels; imputed labels are never scored. Indian event summaries average target errors within state/year/family, states within year, years within family, then families equally. NYC uses the analogous related-mode hierarchy. Ordinary samples and NESO have their declared equal-date/series rules. MAE units remain source-specific. MASE levels across datasets are not direct measures of common task difficulty.

We retain both years, every family, all registered pairwise contrasts and leave-family-out composition tables. We do not use daily-row significance tests or precise annual bootstrap claims from two years. New families, widths, penalties or models are not chosen using evaluated errors.

## 3. Models and representations

Chronos-2 and TimesFM-3 use pinned code/weights, 730 prior target days, fixed seeds, two CPU threads, logged batching and no task fine-tuning. Actual weight bytes and input bytes are checked. Future-target perturbation and solo/batch checks accompany every registered lane; native quantiles and crossings are retained. TimesFM's weight terms include academic noncommercial restrictions. A model release date does not resolve all historical/proxy exposure.

The four local variants are rich Ridge, stable float32 LightGBM, a separately reported minimum-leaf-2 sensitivity, and fixed SARIMAX. Ridge/trees have 27 past-only history features and the same calendar/event covariates; parameters come from pre-evaluation chronological validation with purged future labels and are shared across arms. Monthly local refits differ from frozen foundation inference. Their lag features can draw on observations preceding the 730 labelled training-origin window. Thus total raw-history access and adaptation policies are not identical; model rankings do not identify an architecture-only effect. SARIMAX retains its declared seasonal order, fit/update policy and logged convergence retries. Original and leaf-2 tree budgets are never merged into a chosen winner.

Eight main arms retain 13 ordinary weekday/annual-harmonic channels: calendar, calendar plus five zeros, typed day, pooled day, typed window, pooled window, and typed day/window placebos. Pooled information occupies one channel padded by four zeros; typed information uses five named channels. Windows are ±3-day unions, including subsidiary dates. Fixed seeded weekday shifts avoid the registered true windows, but are bounded active controls, not a randomisation distribution or proof of semantic neutrality. Equal dimensions do not imply equal information, and zero padding is not assumed inert.

Ten pure daily encodings include history-only, a zero channel, ordinal weekday/H, one-hot weekday/H, cyclic weekday/H and five context-mean target encodings. H means any declared named-family anchor, not every public holiday. The hourly paper's hour feature is omitted. Means use only the 730-day input context, with fixed global-context fallback for absent categories; support and future fallback are logged. The unchanged 13-channel calendar reference is added by exact forecast reuse. Channel and information differences are explicit, so a pure encoding versus calendar contrast is not an isolated coding effect. Document-aligned pure encodings use the same corrected cases and ledger as the main analysis.

Eight direct daily comparators include conditional means and no-response/typed/pooled/fixed-partial residual profiles. The latter use a last-four-recorded-same-weekday baseline, a declared triangular ±3-day shape and fixed penalty 2. They are daily adaptations, not an exact replication of Alipour's residual tensors, learned 16-dimensional contrastive representations or past-occurrence prototypes. The primary paper's incomplete implementation details and the lack of a comparable account-level panel remain reproduction gaps.

A separately registered explanatory diagnostic measures recorded support and past family residual heterogeneity within each exact 730-day pre-origin context. Its baseline uses four preceding recorded same-weekday labels outside all named ±3-day windows, with no residual until four exist. Family residual means are scaled by the fixed pre-year weekly scale; their population RMS dispersion requires at least two supported families. Prior-year and current-year support days are separated, and undefined moderators propagate to their cell without dropping forecast errors. Associations round cell variables to eight decimals to define numerical ties. These are observational proxies, not model-learned or causal effects.

## 4. Matched Indian evidence

Tables 1–2 display the prespecified next-day comparison; Figure 1 includes all six variants, both representations and all three leads. Every zero/placebo arm remains in the supplement. Original four-year Chronos seven-channel Diwali deterioration is preserved as a different cohort/definition, not erased by later results.

{main_tables['india_event']}
{main_tables['india_ordinary']}

![Figure 1. Typed minus pooled MASE on the same document-aligned Indian cases. Both signs, all variants and leads are retained.](figures/india_contrasts.png)

The full event contrast range is {lo:.5f} to {hi:.5f}; it describes registered cells rather than a selected gain. Same-case aligned-feature changes at lead 1 range from {claims['aligned_feature_change_min']:.5f} to {claims['aligned_feature_change_max']:.5f}. Separately, replacing the three date/lead cases under original features changes the aggregate by {claims['composition_change_min']:.5f} to {claims['composition_change_max']:.5f}. These are different estimands; an old/new macro-score difference must not be attributed entirely to corrected festival features.

{main_tables['pure']}

Table 3 gives every registered pure encoding at the same lead, without selecting a recommended encoding. Complete lead-7/14, ordinary-day, sparse-support and all-pair tables accompany it. Learned category means are past-only summaries of this dataset, not newly discovered holidays or learned causal event effects.

## 5. Cross-source checks and pooling conditions

{main_tables['mta_event']}

![Figure 2. NYC typed minus pooled MASE. Columns are decision leads 1/7/14, corresponding to actual model steps 8/14/21 under the observation buffer.](figures/mta_contrasts.png)

NYC's complete event contrast range is {mlo:.5f} to {mhi:.5f}. Its local checkpoint already supplied a counterexample to universal pooling gains: SARIMAX typed-day event MASE improves on calendar while pooled-day worsens it, whereas pooling helps the ordinary sample. The full table and Figure 2 now retain both foundation models and every registered contrast. The ordinary table is in Appendix A. Two parking-observance occurrences and two related modes support a bounded robustness check, not universal mechanism identification or a pooled cross-country treatment effect.

{main_tables['neso']}

The NESO comparison keeps ideal and delayed information separate, including all leads and controls. Delay need not monotonically worsen realised error on a small serial sample. It changes the available information and forecast step, not only a calendar label. These ordinary dates supply no additional festival occurrences.

The prespecified simulation has 16 combinations of event frequency, effect heterogeneity, date-error bound and noise, each with 32 seeds: 512 datasets and 2,560 method cells. Five families share a known triangular shape and an exactly removed baseline; noise is stationary AR(1) with coefficient 0.5. No-event, pooled OLS, typed OLS, fixed partial pooling and the unavailable true-effect oracle are all retained. Homogeneous versus heterogeneous effects illustrate the familiar pooling trade-off under that generator. They do not explain foundation-model internals or constitute new real-world replication.

![Figure 3. Synthetic pooling conditions. Intervals quantify Monte Carlo precision over generator seeds, not significance over Indian or NYC events.](figures/simulation.png)

The past-support diagnostic reuses all 6,336 typed/pooled case contrasts, producing 864 family/state/year/lead cells and 1,728 descriptive association records across all variants, representations, leads, moderators and fixed scopes. Table A3 reports the prespecified next-day day-representation support-ratio and response-heterogeneity Spearman diagnostics; both years, all leads and leave-family-out results remain in the supplement. We do not pool models as independent observations or infer a mechanism from a correlation. State/family/year composition and earlier responses from the same event episode can confound these associations. Current-year past responses are legitimate available information, not evidence of forecasting a wholly unseen episode.

## 6. Native uncertainty and reproducibility

Both foundations are compared on unchanged common quantiles 0.1–0.9. Empirical 80% coverage accompanies scaled width, mean marginal pinball and one-interval WIS: (0.5 × median absolute error + 0.1 × 80% interval score)/1.5. Native crossings are retained. Any invalid lower/upper endpoint makes the associated interval aggregate undefined rather than silently dropping cases. Native interval quality and point MASE are different criteria.

{main_tables['native']}

Separate audits reconstruct raw source transformations, primary dates, cohorts, weekly scales, actual inputs, local fit/query arrays, label purging, unchanged-forecast reuse and scalar/hierarchical metrics. NYC additionally reruns all 72 past-only validation fits from independently constructed matrices using the same numerical model libraries; maximum validation-MAE difference is zero. Complete clean-environment reruns of all 5,952 direct-comparator forecasts and 2,560 simulation cells are byte-identical. A complete isolated-environment NYC rerun also reproduces all 4,032 local forecasts as byte-identical prediction files and all 72 validation MAEs exactly, including the declared SciPy 1.18.1 versus 1.17.0 package-drift sensitivity. These are internal software checks on one machine, not independent human review or independent foundation reimplementation.

Earlier numerical failures remain material. A float64 tree reproduction failed before a separately named float32 revision. In the stricter scalar review, 135 forecasts agree within tolerance, but 15 of 45 Delhi lead-1 near-zero feature entries differ after float32 conversion. That discrepancy is not removed by later tests. Failed calendar preflight, interrupted report generation and output-directory-only recovery are retained with their original source and result hashes.

The historical daily CQR study [8] improves some coverage at the cost of wider intervals and some worse longer-lead proper scores. Its cohort differs from festival-feature distributions; it is not presented as their calibration. The original 30 October future freeze contains TimesFM history/calendar, Ridge, the original float64 tree and naive forecasts, with no Chronos or festival-feature arms. Its later interval annex repeats the same TimesFM episode with actual later issuance. Both remain immutable and unscored until the registered collection date. They do not directly confirm the main typed-versus-pooled hypothesis. A separate manual rolling plan is not activated and has issued no forecasts.

{future_text}

## 7. Limits, interpretation and submission gates

The evidence is conditional on source, event definition, representation, horizon and adaptation policy. It cannot establish a new architecture, a causal holiday effect, contamination-free historical testing or a universal best model. India has two main event years; NYC contributes two occurrences from one city; NESO contributes no holiday targets. Weather, supply constraints, transit conditions, coincident observances, source revisions and calendar amendments remain alternative explanations. Local total-history access and fitting policies differ from foundation inference. Row counts and additional models do not create independent events.

The named-date correction demonstrates why checking whether a date is a holiday is weaker than checking which festival the document names. Annual-as-issued alignment is still not final amendment certification. Maharashtra 2022 evidence, shared-main exceptions, first-publication histories and independent human review remain open. Generic encodings and sparse pooling have close prior work; an exact runnable impact-driven comparison or a clearly justified alternative still needs resolution before a novelty claim is defensible.

Indian source-reuse permission is pending after the author manually emailed the curator. Public access and GitHub publication do not imply a grant. The MTA attached Overview explicitly states no data-use limitations, supporting the attributed minimal filtered snapshot; it does not name a standard licence or authorise Indian data. NESO has custom terms whose full body could not be retrieved; raw and complete derived archives remain outside Git. Any journal data-sharing restriction needs an editor-acceptable explanation. No independent reviewer, professor approval, first-vintage reconstruction or future outcome is inferred.

The immediate scientific next step is a genuinely future-issued feature comparison or repeated untouched events with a verified dated calendar and explicit availability/issuance guards, together with external review of the current evidence. It must have a new protocol, actual completion timestamp and pre-outcome public freeze. Current retrieval failures are not grounds to guess 2026 regional dates or silently modify the old future episode. Journal fit, current author instructions, declarations and data/code arrangements require verification; no acceptance or quartile claim follows from these experiments.

## Data, code and author statements

Completed code, protocols, predictions, failures and audit artifacts are retained at https://github.com/Binay-verma/tsaf_project . New rows are distinguished from exact reused forecasts and composed views; no grand total is treated as independent evidence. Local environments, model weights/caches and restricted or unresolved external raw archives are excluded. This draft is a reproducible review document, not a submission or permission certification.

Author: Binay Verma. Affiliation and corresponding email remain blank at the author's request. Funding, conflicts and any real coauthor contributions await factual confirmation. Binay Verma is not the independent reviewer. No outreach or journal submission is made by this build.

## References

[1] Kreusel et al. (2026). On the Encoding of Calendar Features in Covariate-Informed Time Series Foundation Models. arXiv:2610.13676v1. https://arxiv.org/abs/2610.13676

[2] Alipour et al. Impact-driven event covariates for time-series foundation models. Primary page and linked paper have unresolved venue/name metadata; see the retained primary-method audit. https://www.amazon.science/publications/impact-driven-event-covariates-for-time-series-foundation-models

[3] Sanyal, Mitra, McElroy and Roy (2017). Holiday Effects in Indian Manufacturing Series. Census RRS2017-04. https://www.census.gov/library/working-papers/2017/adrm/rrs2017-04.html

[4] Kaljevic, Arzola and Zhang (2026). Benchmarking Time Series Foundation Models for Load Forecasting Under Covariate Uncertainty. arXiv:2610.07232v1. https://arxiv.org/abs/2610.07232

[5] Meyer et al. Rethinking Evaluation in the Era of Time Series Foundation Models: (Un)known Information Leakage Challenges. arXiv:2510.13654v3. https://arxiv.org/abs/2510.13654

[6] Ansari et al. (2025). Chronos-2: From Univariate to Universal Forecasting. arXiv:2510.15821v1. https://arxiv.org/abs/2510.15821

[7] Jain and Sen (2026). TimesFM-3: A zero-shot foundation model for multivariate forecasting. Official model announcement. https://research.google/blog/timesfm-3-a-zero-shot-foundation-model-for-multivariate-forecasting/

[8] Romano, Patterson and Candès (2019). Conformalized Quantile Regression. arXiv:1905.03222v1. https://arxiv.org/abs/1905.03222

## Appendix A. Complete companion comparisons

{appendix}

The supplementary source tables retain all leads, both cohorts, both feature views, every control and every registered pair. Figure and table values are copied from audited aggregates; the build manifest records exact source hashes. The companion audit compares every DOCX table cell with its exported CSV. PDF rendering and visual review are separate presentation checks, not independent scientific review.
'''
    (OUT/'MANUSCRIPT.md').write_text(text)
    metadata = []
    for table in tables:
        path = OUT/'supplement'/(table['name']+'.csv')
        metadata.append(dict(name=table['name'], title=table['title'],
            csv=str(path.relative_to(OUT)), rows=len(table['frame']),
            columns=list(table['frame'].columns), sha256=sha(path)))
    dump(OUT/'tables.json', metadata)
    dump(OUT/'numeric_claims.json', claims)
    docx = OUT/'Integrated_Review_Binay_Verma.docx'
    subprocess.run(['pandoc', str(OUT/'MANUSCRIPT.md'), '-o', str(docx),
        '--resource-path', str(OUT), '--reference-doc',
        str(ROOT/'manuscript/journal_draft_v1/reference.docx')], check=True)
    format_docx(docx)
    assert len(Document(docx).tables) == len(tables) == 9
    inputs = {**{name: sha(ROOT/name) for name in GATES}, **sources,
        'review/clean_mta_v1/comparison.json': sha(clean_mta_path),
        'results/event_support_v1/independent_audit.json': sha(support_audit_path),
        'results/calendar_recenter_recovery_v1/recovery_registration.json': sha(ROOT/'results/calendar_recenter_recovery_v1/recovery_registration.json'),
        **future_inputs,
        'src/build_integrated_review.py': sha(__file__)}
    dump(OUT/'build.json', dict(created_utc=datetime.now(timezone.utc).isoformat(),
        status='Integrated internal review draft; PDF/numeric/visual audit pending',
        tables=len(tables), figures=3, inputs=inputs,
        outputs={str(p.relative_to(OUT)): sha(p) for p in OUT.rglob('*') if p.is_file()}))
    print('Built new integrated MD/DOCX; 9 tables, 3 figures; no old draft changed')


if __name__ == '__main__':
    main()
