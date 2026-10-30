"""Create an editable review manuscript only from completed audited studies."""
from datetime import datetime, timezone
import json
import re
from pathlib import Path
import shutil
import subprocess
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from forecasting import ROOT
from festival_feature_design import sha, dump
from report_multifestival import markdown

OUT = ROOT/'manuscript/journal_draft_v1'
A = ROOT/'results/pooled_festival_v1'
B = ROOT/'results/festival_label_transfer_v2'
C = ROOT/'results/postweight_daily_v1'
MODELS = ['ridge', 'lightgbm', 'lightgbm_capacity2', 'sarimax', 'chronos', 'timesfm']
NAMES = dict(zip(MODELS, ['Rich Ridge', 'Stable LightGBM', 'Tree leaf minimum 2', 'SARIMAX', 'Chronos-2', 'TimesFM-3']))


def polish_markdown(text):
    replacements={'Table1':'Table 1','Table2':'Table 2','Table3':'Table 3','Table4':'Table 4',
        'Figure1':'Figure 1','Figure2':'Figure 2','Figure3':'Figure 3','Figure4':'Figure 4',
        'empirical80%coverage':'empirical 80% coverage','one-intervalWIS':'one-interval WIS',
        'Native80%coverage':'Native 80% coverage','nominal0.8':'nominal 0.8',
        'all135checked':'all 135 checked','maximum7.617e-12GWh':'maximum 7.617e-12 GWh',
        'up to1.421e-14GWh':'up to 1.421e-14 GWh','252predictions':'252 predictions',
        'issued30October2026':'issued 30 October 2026','216eligible rows':'216 eligible rows',
        'for31October–11November':'for 31 October–11 November','45tests':'45 tests',
        'all141training blocks':'all 141 training blocks','Version2 preflights':'Version 2 preflights',
        'nominal80%':'nominal 80%','80%interval':'80% interval'}
    for old,new in replacements.items(): text=text.replace(old,new)
    return text


def format_docx(path):
    doc = Document(path)
    for name in ['Normal', 'Body Text', 'First Paragraph', 'Compact']:
        if name in doc.styles:
            style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(11)
            style.paragraph_format.line_spacing=1.15;style.paragraph_format.space_after=Pt(6)
    for name in ['Heading 1','Heading 2','Heading 3']:
        if name in doc.styles:
            doc.styles[name].font.name='Times New Roman';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    for section in doc.sections:
        section.top_margin=Inches(.8);section.bottom_margin=Inches(.8);section.left_margin=Inches(.85);section.right_margin=Inches(.85)
    for table in doc.tables:
        table.autofit=False
        if 'Table Grid' in doc.styles:
            table.style='Table Grid'
        width=6.8/len(table.columns)
        for column in table.columns: column.width=Inches(width)
        previous=table._tbl.getprevious()
        if previous is not None:
            props=previous.find(qn('w:pPr'))
            if props is None:
                props=OxmlElement('w:pPr');previous.insert(0,props)
            keep=OxmlElement('w:keepNext');props.append(keep)
        for i,row in enumerate(table.rows):
            for cell in row.cells:
                cell.width=Inches(width)
                for paragraph in cell.paragraphs:
                    paragraph.style=doc.styles['Normal'];paragraph.paragraph_format.space_after=Pt(3)
                    paragraph.paragraph_format.line_spacing=1.0
                    paragraph.paragraph_format.keep_with_next=i < len(table.rows)-1
                    paragraph.paragraph_format.keep_together=True
                    for run in paragraph.runs: run.font.size=Pt(9)
    footer=doc.sections[0].footer.paragraphs[0]
    if not footer._p.xpath('.//w:fldSimple'):
        footer.alignment=1
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.save(path)


def main():
    if OUT.exists():
        raise FileExistsError('Version a new manuscript; preserve review drafts')
    for path in [A/'independent_audit.json', A/'inputs_fits_audit.json', A/'probabilistic_audit.json', B/'independent_audit.json', C/'independent_audit.json']:
        assert json.loads(path.read_text())['status'] == 'pass'
    OUT.mkdir(parents=True)
    figures = OUT/'figures'; figures.mkdir()
    for cohort in ['event', 'non_event']:
        shutil.copyfile(A/f'{cohort}_typed_minus_pooled5.png', figures/f'{cohort}_typed_minus_pooled5.png')
    event = pd.read_csv(A/'event/primary_metrics_balanced.csv')
    event1 = event[event.horizon.eq(1)].pivot(index='model', columns='arm', values='MASE').reindex(MODELS)
    table_a = event1[['calendar', 'pooled_day1', 'pooled_day5', 'day5']].rename(index=NAMES).reset_index()
    table_a.columns = ['Model', 'Calendar', 'Pooled day: 1ch', 'Pooled day: 5ch', 'Typed day: 5ch']
    daily = pd.read_csv(C/'metrics_macro.csv')
    configs = json.loads((C/'protocol.json').read_text())['configurations']
    table_c = daily.pivot(index='configuration', columns='horizon', values='MASE').reindex(configs)
    table_c.columns = ['Lead '+str(h) for h in table_c.columns]
    fig, ax = plt.subplots(figsize=(10, 5.5))
    for config in configs:
        a = daily[daily.configuration.eq(config)].sort_values('horizon')
        ax.plot(a.horizon, a.MASE, marker='o', linewidth=1.5, label=config.replace('_', ' '))
    ax.set_xticks([1,7,14]); ax.set_xlabel('Individual forecast lead (days)'); ax.set_ylabel('Equal-state MASE (lower better)')
    ax.set_title('Retrospective daily panel: 44 matched target dates per state/lead')
    ax.grid(alpha=.2); ax.legend(fontsize=8, ncol=2)
    fig.tight_layout(); fig.savefig(figures/'daily_point_errors.png', dpi=190); plt.close(fig)
    transfer = pd.read_csv(B/'metrics_balanced.csv')
    t = transfer[transfer.cohort.eq('event') & transfer.horizon.eq(1)].copy()
    t[['feature', 'policy']] = t.arm.str.split('__', expand=True)
    table_b = t.pivot(index=['model','policy'], columns='feature', values='MASE')[['calendar','pooled_day5','typed_day5']].reset_index()
    native = pd.read_csv(C/'native_uncertainty_macro.csv')
    native1 = native[native.horizon.eq(1)][['configuration', 'scaled_mean_pinball', 'coverage80', 'scaled_width80', 'scaled_WIS80_median']]
    fig, axes = plt.subplots(1,3,figsize=(12,4.2),sharey=True)
    colors = dict(zip(configs[:4], ['#315a92','#6997c9','#a15c24','#d6a16d']))
    for ax,h in zip(axes,[1,7,14]):
        for r in native[native.horizon.eq(h)].itertuples(index=False):
            ax.scatter(r.scaled_width80, r.coverage80, color=colors[r.configuration], s=65, label=r.configuration.replace('_',' '))
        ax.axhline(.8,color='black',linestyle='--',linewidth=1)
        ax.set_ylim(0,1.03);ax.set_title('Lead '+str(h));ax.set_xlabel('Mean scaled 80% width');ax.grid(alpha=.2)
    axes[0].set_ylabel('Empirical native 80% coverage')
    handles,labels=axes[0].get_legend_handles_labels();fig.legend(handles,labels,loc='lower center',ncol=4,fontsize=8)
    fig.tight_layout(rect=[0,.12,1,1]);fig.savefig(figures/'native_coverage_width.png',dpi=190);plt.close(fig)
    a_con = pd.read_csv(A/'event/contrasts_balanced.csv')
    typed_day = a_con[a_con.arm.eq('day5') & a_con.reference.eq('pooled_day5')]
    counts = dict(typed_lower=int(typed_day.MASE_change.lt(-1e-9).sum()), pooled_lower=int(typed_day.MASE_change.gt(1e-9).sum()), numerical_ties=int(typed_day.MASE_change.abs().le(1e-9).sum()))
    b_con = pd.read_csv(B/'contrasts_balanced.csv')
    transfer_pool = b_con[b_con.cohort.eq('event') & b_con.arm.eq('pooled_day5__exclude_family') & b_con.reference.eq('calendar__exclude_family')]
    cover = json.loads((ROOT/'review/used_calendar_v2/coverage.json').read_text())
    title = 'Festival representation and adaptation in forecasting Indian supplied electricity: matched controls, transfer and reproducibility'
    lines = ['# '+title, '', '**Working manuscript for author and professor review — 30 October 2026. Not submitted or submission-ready.**', '',
        'Author: Binay Verma. Affiliation, corresponding-author contact, coauthor/contribution decisions and funding/conflict declarations require author confirmation. No professor authorship or independent sign-off is assumed.', '',
        '## Abstract', '',
        f'Moving festivals create forecasting challenges that ordinary weekday and annual-seasonality features may not capture. We study whether explicit festival representations help predict daily supplied electricity in Delhi, Karnataka and Maharashtra, and whether separate family identities add value beyond a pooled event flag. A named 2022–2025 calendar ledger supports five festival families. Fixed historical anchor and non-event cohorts compare four classical/machine-learning variants and two pinned foundation models under equal-case and feature-count controls. In the next-day anchor cohort, rich Ridge MASE changes from {event1.loc["ridge","calendar"]:.3f} with ordinary calendar inputs to {event1.loc["ridge","day5"]:.3f} with separate day flags and {event1.loc["ridge","pooled_day5"]:.3f} with a five-channel padded pooled flag. Separate-versus-pooled effects differ across model, lead and representation. A further fixed-budget experiment removes own-family fitting labels with matched training-row controls; earlier target history remains allowed. A broader daily panel uses 44 target dates after both audited weight uploads and evaluates native quantiles as well as point errors. Independent software implementations reveal that forecast agreement can coexist with byte-level feature discrepancies. These findings support careful representation, adaptation and provenance controls rather than universal model or festival-enhancement claims. Previously inspected data, correlated events, unresolved source terms, source revisions, calendar amendments and incomplete independent review limit confirmatory interpretation.', '',
        '**Keywords:** electricity forecasting; moving holidays; time-series foundation models; calendar covariates; temporal evaluation; reproducibility.', '',
        '## 1. Introduction', '',
        'Electricity supplied around a moving holiday can differ from an ordinary weekday because economic activity, institutional closures, weather and supply conditions change together. A holiday label is therefore potentially useful predictive information, but it is not a causal measure of demand reduction. The problem is particularly delicate when a nominally shared festival has different regional dates, office/bank definitions or subsidiary observances.', '',
        'Time-series foundation models can accept future-known covariates, but apparent improvements can conflate event dates, representation, channel count, statistical sharing and adaptation policy. Separate festival channels may fit distinct responses while receiving only one or two historical positive examples. A pooled flag has more examples but assumes greater sharing. Extra channels can alter foundation-model forecasts even when they contain zeros. These possibilities require direct controls rather than a comparison of errors during festivals alone.', '',
        'Calendar encodings and learned event covariates are established prior work [1,2]. Load forecasting under covariate uncertainty [3], foundation-model evaluation leakage [4] and Indian moving-holiday regression [8] also precede this study. We make no first-ever or new-architecture claim. The candidate empirical contribution is a documented Indian regional-calendar evaluation that separates added-date information from representation and padding, examines supervised family-label transfer, and retains numerical and source-vintage failures.', '',
        'We ask four questions. RQ1: do fixed explicit festival features improve matched historical errors beyond ordinary calendar, zero and placebo controls? RQ2: do separate family channels help beyond a pooled event flag at the same channel budget? RQ3: can pooled information help when direct fitting labels around one family are withheld, compared with removing the same number of ordinary labels? RQ4: how do broader post-weight daily errors, native uncertainty and reproduction limits qualify the conclusions? All new-method evaluations on already inspected historical data remain exploratory.', '',
        '## 2. Data and provenance', '',
        'The target is recorded daily electricity energy met/supplied, in GWh (equivalently MU), rather than latent unconstrained demand or hourly peak power. We use a frozen Robbie Andrew compilation of GRID-INDIA/POSOCO state data ending 28 October 2026. It is kept separate from an older corrected Hunt/Bloomfield CC BY 4.0 source. Shared-source national derivatives are not independent replication.', '',
        'Missing historical inputs are filled causally using a preceding weekly value, then the preceding available value where necessary. Recorded positive labels alone are scored; missing labels are not imputed for evaluation. A separate provenance study verifies the frozen Git source bytes and older archive payload digests. Observed historical revisions reach 61.5 GWh, and two fixed publication-vintage scenarios show context lag. Sparse captures do not reconstruct every historical release. Ideal origin-1 availability must therefore be distinguished from a real-time operational trial.', '',
        'The named feature ledger contains 58 anchors in 48 family/state/year groups across 2022–2025. Holi is evaluated in Delhi/Maharashtra; Ganesh Chaturthi and Ugadi/Gudi Padwa in Karnataka/Maharashtra; Dussehra and Diwali in all three states. State scope is part of the definition. Karnataka lists Ganesh on 18 September 2023 whereas Maharashtra lists 19 September; Maharashtra Holi is 7 March 2023. Annual-as-issued notices, library dates, restricted leave, bank holidays and religious observances are not interchangeable.', '',
        f'The current used-anchor crosswalk covers all 58 rows and 12 state-years: {cover["annual_schedule_date_matches"]} dates match the audited state annual schedule, while {cover["shared_main_without_state_schedule_match"]} declared shared-main dates do not. Retained state PDFs cover {cover["anchors_with_retained_state_pdf"]} anchor rows and previously strengthened notice dates cover {cover["anchors_with_previously_strengthened_notice_date"]}. Shared-main Maharashtra Diwali anchors are not relabelled as state Lakshmi Pujan dates. None of the 58 amendment chains or human sign-offs is certified. Known notice dates are not automatically first historical public-availability times.', '',
        '## 3. Experimental design', '',
        '### 3.1 Temporal rules and metrics', '',
        'An origin is the first unseen date; individual lead h targets origin+h−1. Every input target observation precedes the origin. Direct local-model fitting purges labels beyond the refit cutoff, and scaler fitting uses only retained rows. Parameters are selected on chronological pre-year validation when that policy is declared; no evaluated error selects a new representation, case, budget or deployment winner.', '',
        'MAE is reported in GWh by state. MASE divides absolute error by the recorded mean absolute weekly difference in the preceding annual 730-calendar-day reference window. This reference stays fixed across arms/refits within year. Event summaries average states within year, years within family, and families equally. Daily summaries average states equally. Relative MAE effects are computed within state/year before aggregation; they are not ratios of macro MASE. Near-zero denominators are not stabilized by an arbitrary epsilon. More related forecast rows do not create more independent event years.', '',
        '### 3.2 Models and adaptation', '',
        'Classical/machine-learning models include rich-history Ridge, stable-interface LightGBM, a separately fixed leaf-minimum-2 sensitivity, fixed-order annual SARIMAX, and earlier weekly ETS/persistence comparators. Rich regressions use 27 causal lag/rolling/trend/annual-history predictors and 13 ordinary weekday/Fourier features. They refit monthly using the original pre-year-selected parameters in the feature intervention. Local feature warm-up can extend before the 730-day fitting-label window; adaptation and memory budgets are not architecture-identical.', '',
        'Chronos-2 and TimesFM-3 perform inference without task fine-tuning, with 730 past targets, 14 forecast days, seed 42, CPU batch 4 and two threads. Independent-task attention is enforced for Chronos; TimesFM symmetric averaging is disabled. Actual weight bytes and code versions are audited: Chronos weights SHA-256 ddcda3c7… first uploaded 30 October 2025; TimesFM weights a7592b0a… first uploaded 24 August 2026. Historical/proxy exposure remains unresolved. TimesFM model terms permit the noncommercial academic setup; weights are not redistributed.', '',
        '### 3.3 Added features and pooled comparison', '',
        'The first intervention uses the same 72 anchor/lead cases across six variants: 24 family/state/year targets at leads 1/7/14. Its arms are ordinary calendar, five zeros, five family day flags, five family ±3-day window flags, and same-size seeded shifted day/window placebos. Wider classical windows contain 528 cases; their scores are not compared with foundation anchor-only scores. The matched non-event cohort has another 72 cases, selected by fixed weekday-preserving 28–42-day shifts excluding annual/registered holiday windows.', '',
        'The pooled follow-up adds pooled day/window flags as either one channel or one active union channel plus four exact zeros, together with a one-zero control. Calendar/typed predictions are reused from hash-checked parents. It produces 4,320 new forecasts and 2,592 reused forecasts across the two distinct cohorts. Typed versus pooled5 fixes dimensions but changes signal allocation and statistical sharing; it is not causal identification of identity. Pooled1 versus pooled5 checks padding sensitivity. All channel names, dates, contexts and parameters are retained.', '',
        '### 3.4 Supervised family-label transfer', '',
        'A separate study fixes Ridge alpha=10 and two 7-leaf, 200-estimator, lambda-10 tree budgets with leaf minima 20/2. It does not reuse hyperparameters whose validation had seen the held family. Calendar, typed-day5 and pooled-day5 are crossed with all rows, removal of every direct own-family ±3-day fitting label, and removal of the same number of eligible ordinary labels. Model/scaler fitting uses retained rows only.', '',
        'Control dates are matched deterministically by weekday, year, training-window availability and holiday exclusion, using fixed ±28, ±35, …, ±84-day shifts. The first version correctly aborted when a 28–42-day match was infeasible; its partial outputs and failure are retained. Version2 preflights all 141 training blocks before model execution. Larger shifts weaken seasonal comparability. Crucially, past electricity history and annual lags can still contain earlier family responses: this is supervised label exclusion with available target history, not fully unseen-event or history-free zero-shot transfer. Three budgets × nine arms × 144 cases produce 3,888 forecasts.', '',
        '### 3.5 Daily post-weight panel and native uncertainty', '',
        'A broader panel uses 44 identical target dates, 15 September–28 October 2026, for three states and three leads: 396 cases per configuration. All origins are 2 September or later, after both audited weight uploads and the documented TimesFM announcement. Eight configurations compare history/calendar in both foundation models, original pre-2026-selected monthly rich-calendar Ridge/tree, last value and weekly repeat. This yields 3,168 forecasts and 684 foundation context/mode tasks. No 2026 festival calendar or future weather is supplied. Previously inspected outcomes remain exploratory despite post-weight chronology.', '',
        'Native marginal quantiles 0.1–0.9 are scored by mean scaled pinball. Central 80% intervals are evaluated using coverage, lower/upper misses, scaled width and interval score. A one-interval WIS combines the native median absolute error and the 80% interval score; it is not a full-grid WIS. No quantile sorting, clipping or calibration is applied. Invalid endpoints remain explicit and propagate undefined interval summaries. These native intervals differ from an earlier past-residual nominal-90% interval study.', '',
        '## 4. Results', '',
        '### 4.1 Explicit and pooled festival features', '',
        'Table1 reports all six variants on identical next-day anchor cases. It is an error table, not accuracy percentages. All other leads, windows, zeros/placebos, states and years remain in the versioned reports.', '',
        '**Table1. Next-day anchor MASE: ordinary, pooled and typed day features.**', '', markdown(table_a.round(5)), '',
        f'Across the 18 model/lead cells in the equal-five-channel typed-day versus pooled-day contrast, typed errors are lower in {counts["typed_lower"]}, pooled errors lower in {counts["pooled_lower"]}, with {counts["numerical_ties"]} ties within 1e-9 MASE. These counts are descriptive and correlated, not independent trials or a selection rule. Figure1 shows both day and window contrasts at every lead; Figure2 shows matched non-event effects. Separate family channels are not assumed to dominate pooling.', '',
        '![Figure1. Event-anchor typed minus pooled5 MASE, all variants/leads. Negative favours typed; positive favours pooled.](figures/event_typed_minus_pooled5.png)', '',
        '![Figure2. The same contrast on the separate matched non-event cohort. This is not exhaustive all-day performance or a causal event/control comparison.](figures/non_event_typed_minus_pooled5.png)', '',
        'Earlier controls remain essential: the original four-year seven-channel Diwali enhancement worsens Chronos mean day-1 error, while later differently scoped day-feature gains can be positive. Zero and placebo channels change some foundation forecasts without equivalent event information. Broad windows and spring-family comparisons can worsen errors. No earlier negative result is replaced by the newer anchor cohort.', '',
        '### 4.2 Supervised transfer and matched data-removal controls', '',
        '**Table2. Event-anchor next-day MASE under fixed budgets and row policies.**', '', markdown(table_b.round(5)), '',
        'The primary interpretation compares pooled and calendar inputs within the family-excluded policy, then examines family exclusion versus equally sized control removal for the same feature. Comparing only with the full-data model would confound the kind of lost labels with training-set size. Every own-family typed indicator has zero active fitting rows after exclusion; historical target exposure remains available through lag/rolling features. All-lead and family/state/year contrasts, distance inventory and retained-label hashes are provided in the transfer report. For Ridge under family-label exclusion, pooled MASE worsens at leads 7/14 (1.33987/1.52931 versus calendar 1.18662/1.21157), despite its next-day gain. The leaf-minimum-2 tree has lower aggregate errors at all three leads in this bounded cohort; the leaf-minimum-20 tree ignores the sparse indicators. These results cannot establish foundation-model unseen-event transfer.', '',
        '### 4.3 Daily comparison after actual weight uploads', '',
        '**Table3. Equal-state daily MASE on the 44-target-date panel.**', '', markdown(table_c.round(5).reset_index()), '',
        '![Figure3. All eight daily configurations and individual leads. Lines connect summary points, not a continuous forecast trajectory.](figures/daily_point_errors.png)', '',
        'TimesFM history-only has the lowest macro MASE among the eight configurations in this panel: 0.54620, 1.16303 and 1.30856 at leads 1, 7 and 14. These errors provide a broader two-model comparison than the small festival anchor cohorts. They are conditional on fixed fitting policies and this late-2026 period, not proof of universal superiority or full-year ranking. Post-weight dates support direct temporal exposure evidence, conditional on correct source observation dates; retrospective inspection and source publication lag remain separate issues.', '',
        '### 4.4 Native probability quality and numerical reproduction', '',
        '**Table4. Next-day native uncertainty: common pinball, empirical80%coverage, scaled width and one-intervalWIS.**', '', markdown(native1.round(5)), '',
        '![Figure4. Native80%coverage versus scaled interval width at each lead. The dashed line is nominal0.8, not a guarantee or fitted calibration target.](figures/native_coverage_width.png)', '',
        'Coverage must be read with width and proper scores. Point forecast improvement alone does not imply a calibrated or ordered predictive distribution. In the event-anchor cohort, next-day event-balanced coverage of nominal 80% intervals is 0.100 for calendar-only Chronos and 0.16667 for calendar-only TimesFM; TimesFM pooled-day5 improves coverage to 0.500 but remains below nominal. These are weighted historical proportions, not independent binomial trials. The pooled comparison retains 14 all-native adjacent crossing pairs (9 new, 5 reused) for Chronos and none for TimesFM; its common 0.1–0.9 grid has no crossings. All native adjacent crossings are inventoried and retained. The panel has correlated days/states and overlapping horizons, so nominal coverage is descriptive rather than a reliable calibration certification.', '',
        'Separate internal software audits reconstruct labels, denominators, complete cases, aggregation, paired changes, actual foundation input bytes, fitting labels/cutoffs and probability arithmetic. The bounded clean scalar review reproduces all135checked pooled regression forecasts within tolerance, maximum7.617e-12GWh, while its strict overall status remains discrepancies_retained. One near-zero Delhi trend entry per checked block differs after float32 conversion by up to1.421e-14GWh. The earlier float64 tree study also retains larger forecast discrepancies. Agreement on checked forecasts does not imply universal byte identity or independent human review.', '',
        '## 5. Discussion', '',
        'The study separates several questions that a simple holiday benchmark can collapse. Measuring error during an event does not test incremental event information. Adding family channels does not demonstrate that identities beat pooled dates. Holding dimensions fixed does not equalize statistical sharing. Withholding direct labels does not erase historical event responses from causal target features. Post-weight chronology does not turn retrospectively inspected data into untouched evaluation. Each distinction changes what a positive numerical result can support.', '',
        'For practitioners, the implication is to compare event representations under matched cases and declared adaptation budgets, include ordinary-day costs and report sparse support. Pooled learning can be a useful comparator when separate events have few past examples, but its value must be measured rather than assumed. Calendar provenance should preserve administrative/religious scope and subsidiary anchors; an office schedule is not a direct measure of statewide activity.', '',
        'For evaluation, native uncertainty, numerical feature contracts and source vintages need explicit audits. Minute floating-point differences can matter to trees, and zero channels can matter to foundation models. Source revisions and reporting delays can alter the idealized history available at an origin. These are empirical reproducibility and availability questions, not reasons to suppress failures or revise outcomes for more attractive tables.', '',
        '## 6. Limitations and confirmation plan', '',
        'The current evidence uses one source lineage, three related states, five scoped families and only two primary historical event years. Nearby dates, leads and states are correlated. All historical outcomes have been inspected, so the new-method experiments remain development/exploratory. The placebo/control constructions are bounded and not exchangeable causal assignments. Model/source proxy exposure, exact historical publication times, curator revisions/imputation and full calendar amendments remain unresolved.', '',
        'An existing prospective freeze contains252predictions issued30October2026, with216eligible rows for31October–11November; earlier two target dates are explicitly excluded. The original Git blobs remain unchanged. No prospective outcomes are available here and the auto-scoring workflow remains deliberately paused. One eventual shared-source episode will still not replace independent-source or broader independent-event confirmation.', '',
        'The next confirmation requires a separately frozen protocol, independent or genuinely untouched observations with model-specific chronology, source-reuse clarification, complete used-date amendment review and independent scientific reproduction. The professor/author must verify contribution, bibliography, affiliation, declarations and venue suitability. A Q1 aspiration does not change these evidence requirements. No submission or acceptance is claimed.', '',
        '## 7. Conclusion', '',
        'Calendar/date representation, pooling, adaptation and numerical/source provenance must be distinguished when assessing foundation forecasts of supplied electricity around Indian festivals. The controlled comparisons provide bounded positive and negative findings while retaining earlier results. They support an empirical evaluation and failure-analysis framing; a new architecture, universal holiday benefit and confirmatory generalisation are not established.', '',
        '## Data, code and review statements', '',
        'Code, versioned predictions, protocols, failures and handoffs are synchronized to https://github.com/Binay-verma/tsaf_project under the author’s explicit instruction. Upstream source-reuse permission and processing clarification remain pending after the author’s manual email; public availability does not grant a new licence. Model weights and caches are excluded from the repository; pinned restoration/hash instructions are retained. Any journal data-sharing exception requires actual editorial acceptance and is not assumed.', '',
        'All automated/separate-implementation checks are internal software evidence. No independent human scientific/calendar approval has occurred. Author contributions, competing interests, funding, acknowledgements, affiliation and corresponding-author details require author confirmation. No human-subject claim or declaration is fabricated. No other person has been contacted on the author’s behalf during this work.', '',
        '## References and artifact map', '',
        'Bibliographic details and venue status require final author verification. Preprints and software are not represented as journal publications.', '',
        '1. Kreusel et al. (2026). On the Encoding of Calendar Features in Covariate-Informed Time Series Foundation Models. arXiv:2610.13676. https://arxiv.org/abs/2610.13676', '',
        '2. Alipour et al. Impact-driven event covariates for time-series foundation models. Amazon Science primary page; final author/venue metadata unresolved. https://www.amazon.science/publications/impact-driven-event-covariates-for-time-series-foundation-models', '',
        '3. Kaljevic, Arzola and Zhang (2026). Benchmarking Time Series Foundation Models for Load Forecasting Under Covariate Uncertainty. arXiv:2610.07232. https://arxiv.org/abs/2610.07232', '',
        '4. Meyer et al. Rethinking Evaluation in the Era of Time Series Foundation Models: (Un)known Information Leakage Challenges. arXiv:2510.13654v3. https://arxiv.org/abs/2510.13654', '',
        '5. Ansari et al. (2025). Chronos-2: From Univariate to Universal Forecasting. arXiv:2510.15821. https://arxiv.org/abs/2510.15821', '',
        '6. Shchur et al. fev-bench: A Realistic Benchmark for Time Series Forecasting. arXiv:2509.26468. https://arxiv.org/abs/2509.26468', '',
        '7. Aksu et al. (2024). GIFT-Eval: A Benchmark for General Time Series Forecasting Model Evaluation. arXiv:2410.10393. https://arxiv.org/abs/2410.10393', '',
        '8. Sanyal, Mitra, McElroy and Roy (2017). Holiday Effects in Indian Manufacturing Series. Census RRS2017-04. https://www.census.gov/library/working-papers/2017/adrm/rrs2017-04.html', '',
        '9. Google Research (2026). TimesFM-3 official announcement and pinned model/code. https://research.google/blog/timesfm-3-a-zero-shot-foundation-model-for-multivariate-forecasting/', '',
        '10. Hyndman and Athanasopoulos. Forecasting: Principles and Practice. https://otexts.com/fpp3/toolbox.html', '',
        '11. Source/model/calendar identities and primary URLs: DATA_SOURCES.md, REGIONAL_CALENDAR_AUDIT.md, PRETRAINING_AUDIT.md and results/second_model_audit_v1/REPORT.md.', '',
        '12. Current experiment protocols/results: results/festival_features_v1, festival_feature_spillover_v2, pooled_festival_v1, festival_label_transfer_v2 and postweight_daily_v1. Earlier failed prototypes, original expanded findings, selected parameters and prospective freeze remain versioned. PROTOCOL_INDEX.md provides the complete experiment registry.', '',
        '## Appendix: exact result sources', '',
        'Table1: pooled_festival_v1/event/primary_metrics_balanced.csv. Figures1/2: contrasts_balanced.csv with equal state/year/family weights. Table2: festival_label_transfer_v2/metrics_balanced.csv. Table3/Figure3: postweight_daily_v1/metrics_macro.csv. Table4/Figure4: postweight_daily_v1/native_uncertainty_macro.csv. All full-arm/state/year/lead tables and native outputs accompany the manuscript; no row-count-based significance claim or final-model selection is made.']
    rendered=polish_markdown('\n'.join(lines)+'\n')
    rendered=re.sub(r'!\[([^\]]+)\]\((figures/[^)]+)\)', lambda m: m.group(0)+'\n\n*'+m.group(1)+'*', rendered)
    md = OUT/'MANUSCRIPT.md'; md.write_text(rendered)
    with (OUT/'reference.docx').open('wb') as stream:
        subprocess.run(['pandoc','--print-default-data-file','reference.docx'],stdout=stream,check=True)
    template = Document(OUT/'reference.docx')
    normal = template.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.15; normal.paragraph_format.space_after=Pt(6)
    for section in template.sections:
        section.top_margin=Inches(.8);section.bottom_margin=Inches(.8);section.left_margin=Inches(.85);section.right_margin=Inches(.85)
    template.save(OUT/'reference.docx')
    subprocess.run(['pandoc', str(md), '--from=gfm', '--to=docx', '--standalone', '--resource-path='+str(OUT),
        '--reference-doc='+str(OUT/'reference.docx'), '-o', str(OUT/'Manuscript_Binay_Verma.docx')], check=True, cwd=ROOT)
    format_docx(OUT/'Manuscript_Binay_Verma.docx')
    sources=[A/'event/primary_metrics_balanced.csv',A/'event/contrasts_balanced.csv',B/'metrics_balanced.csv',C/'metrics_macro.csv',C/'native_uncertainty_macro.csv']
    dump(OUT/'build.json',dict(status='editable_draft_created_pdf_QA_pending',created_utc=datetime.now(timezone.utc).isoformat(),
        source_hashes={str(f.relative_to(ROOT)):sha(f) for f in sources},manuscript_sha256=sha(md),
        counts=counts,transfer_pool_contrasts=transfer_pool.to_dict('records'),
        limits='No human review, permission, untouched confirmation, venue quartile or acceptance fabricated'))
    print('Created review manuscript MD/DOCX with4figures; PDF conversion/visual QA still required')


if __name__ == '__main__': main()
