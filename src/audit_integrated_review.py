"""Independently match integrated tables/text claims to audited source rows.

Uses standard-library CSV/scalar grouping and DOCX parsing rather than the
builder's pivots or formatting helpers. This is an internal reporting audit.
"""
import csv
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'manuscript/integrated_review_v3'
BASE = ROOT/'results/calendar_recenter_recovery_v1/analysis'
MODELS = ['ridge', 'sarimax', 'lightgbm', 'lightgbm_capacity2', 'chronos', 'timesfm']
LABELS = {'ridge': 'Rich Ridge', 'sarimax': 'SARIMAX', 'lightgbm': 'Stable LightGBM',
          'lightgbm_capacity2': 'Tree leaf minimum 2', 'chronos': 'Chronos-2', 'timesfm': 'TimesFM-3'}
ARMS = ['calendar', 'day5', 'pooled_day5', 'window5', 'pooled_window5']
ARM_LABELS = ['Calendar', 'Typed day', 'Pooled day', 'Typed window', 'Pooled window']


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read(path):
    with Path(path).open(newline='') as stream:
        return list(csv.DictReader(stream))


def fmt(value):
    number = float(value) if value != '' else math.nan
    return 'undefined' if math.isnan(number) else f'{number:.5f}'


def select(rows, **query):
    matches = [row for row in rows if all(row[k] == str(v) for k, v in query.items())]
    assert len(matches) == 1, (query, len(matches))
    return matches[0]


def scores(rows, *, version=False, lead_field='horizon'):
    query = {'version': 'document_aligned'} if version else {}
    return [['Model']+ARM_LABELS]+[
        [LABELS[model]]+[fmt(select(rows, model=model, arm=arm,
            **{lead_field: 1}, **query)['MASE']) for arm in ARMS] for model in MODELS]


def main():
    if (OUT/'independent_reporting_audit.json').exists():
        raise FileExistsError('Preserve previous reporting audit')
    build = json.loads((OUT/'build.json').read_text())
    for name, expected in build['inputs'].items():
        assert sha(ROOT/name) == expected, name
    for name, expected in build['outputs'].items():
        assert sha(OUT/name) == expected, name
    source_index = json.loads((OUT/'supplement/source_index.json').read_text())
    for name, info in source_index.items():
        assert build['inputs'][info['original']] == info['sha256']
        assert sha(ROOT/info['original']) == sha(OUT/name) == info['sha256']
    expected = {}
    for cohort, name in [('event_window', 'india_event'), ('ordinary', 'india_ordinary')]:
        expected[name] = scores(read(BASE/'breadth'/cohort/'metrics_balanced.csv'), version=True)
    pure = read(BASE/'encoding/event_window/metrics_balanced.csv')
    order = ['history', 'zero1', 'ordinal', 'onehot', 'cyclic', 'te_weekday',
             'te_sunday', 'te_holiday', 'te_joint', 'te_separate', 'calendar']
    expected['pure'] = [['Encoding', 'Chronos-2', 'TimesFM-3']]+[
        [arm]+[fmt(select(pure, model=model, arm=arm, horizon=1,
                         version='document_aligned')['MASE']) for model in ['chronos', 'timesfm']]
        for arm in order]
    mta = ROOT/'results/mta_holiday_v1/analysis'
    expected['mta_event'] = scores(read(mta/'event_window/metrics_balanced.csv'), lead_field='decision_lead')
    expected['mta_ordinary'] = scores(read(mta/'ordinary/metrics_balanced.csv'), lead_field='decision_lead')
    neso = read(ROOT/'results/neso_replication_v1/analysis/metrics.csv')
    expected['neso'] = [['Model', 'Ideal', 'Lag 21']]+[
        [LABELS[model]]+[fmt(select(neso, model=model, arm='calendar', decision_lead=1,
                                  scenario=scenario)['MASE']) for scenario in ['ideal', 'lag21']]
        for model in MODELS]
    native = read(BASE/'breadth/event_window/native_uncertainty_balanced.csv')
    expected['native'] = [['Model', 'Arm', 'Coverage 80%', 'Scaled width', 'Scaled WIS80']]
    for row in native:
        if row['version'] == 'document_aligned' and row['horizon'] == '1' and row['arm'] in ARMS:
            expected['native'].append([LABELS[row['model']], ARM_LABELS[ARMS.index(row['arm'])]]+
                [fmt(row[key]) for key in ['coverage80', 'scaled_width80', 'scaled_WIS80_median']])
    comparator = read(ROOT/'results/event_comparators_v1/metrics_balanced.csv')
    expected['response_comparators'] = [['Method', 'Lead 1', 'Lead 7', 'Lead 14']]+[
        [method]+[fmt(select(comparator, cohort='event_window', method=method,
                            horizon=lead)['MASE']) for lead in [1, 7, 14]]
        for method in ['response_none', 'response_partial', 'response_pooled', 'response_typed']]
    support = read(ROOT/'results/event_support_v1/associations.csv')
    expected['support_diagnostics'] = [['Model', 'Support rho', 'Support n', 'Response rho', 'Response n']]
    for model in MODELS:
        values = []
        for metric in ['own_support_ratio', 'heterogeneity_scaled']:
            row = select(support, model=model, representation='day', horizon=1, scope='all', moderator=metric)
            values += [fmt(row['spearman']), str(int(row['defined_cells']))]
        expected['support_diagnostics'].append([LABELS[model]]+values)
    doc = Document(OUT/'Integrated_Review_Binay_Verma.docx')
    metadata = json.loads((OUT/'tables.json').read_text())
    assert len(doc.tables) == len(metadata) == len(expected) == 9
    checked = 0
    for info, table in zip(metadata, doc.tables):
        truth = expected[info['name']]
        actual = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        assert actual == truth, (info['name'], actual, truth)
        assert sha(OUT/info['csv']) == info['sha256']
        rows = read(OUT/info['csv'])
        assert list(rows[0]) == truth[0]
        for record, wanted in zip(rows, truth[1:]):
            for key, value in zip(truth[0], wanted):
                if key in ['Support n', 'Response n']:
                    assert str(int(record[key])) == value
                elif value in ['undefined'] or key not in ['Model', 'Method', 'Arm', 'Encoding']:
                    assert fmt(record[key]) == value, (info['name'], key)
                else:
                    assert record[key] == value
        assert len(rows) == len(truth)-1
        checked += sum(len(row) for row in truth[1:])
    # Reconstruct all scalar narrative claims, not merely their presence in text.
    claims = json.loads((OUT/'numeric_claims.json').read_text())
    for source, stem, version, lead_field in [
        (BASE/'breadth/event_window/metrics_balanced.csv', 'india', True, 'horizon'),
        (mta/'event_window/metrics_balanced.csv', 'mta', False, 'decision_lead')]:
        rows = read(source)
        values = []
        for model in MODELS:
            for arm, reference in [('day5', 'pooled_day5'), ('window5', 'pooled_window5')]:
                for lead in [1, 7, 14]:
                    query = dict(model=model, **{lead_field: lead})
                    if version:
                        query['version'] = 'document_aligned'
                    a = float(select(rows, arm=arm, **query)['MASE'])
                    b = float(select(rows, arm=reference, **query)['MASE'])
                    values.append(a-b)
        assert math.isclose(min(values), claims[stem+'_min'], abs_tol=1e-12)
        assert math.isclose(max(values), claims[stem+'_max'], abs_tol=1e-12)
    for filename, column, stem in [
        ('changes_balanced.csv', 'MASE_change', 'aligned_feature_change'),
        ('composition_changes.csv', 'MASE_composition_change', 'composition_change')]:
        values = [float(row[column]) for row in read(BASE/'breadth/event_window'/filename) if row['horizon'] == '1']
        assert math.isclose(min(values), claims[stem+'_min'], abs_tol=1e-12)
        assert math.isclose(max(values), claims[stem+'_max'], abs_tol=1e-12)
    text = (OUT/'MANUSCRIPT.md').read_text()
    for value in claims.values():
        assert f'{value:.5f}' in text
    assert text.count('![Figure ') == 3
    assert 'no Chronos or festival-feature arms' in text
    assert 'not independent human review' in text
    audit = dict(status='pass', audited_utc=datetime.now(timezone.utc).isoformat(),
        scope='Internal independent CSV/scalar reconstruction of every integrated DOCX table cell and eight narrative claims; no independent human scientific/visual approval.',
        tables=9, data_cells_checked=checked, numeric_claims=8, figures=3,
        source_files_checked=len(build['inputs']), docx_sha256=sha(OUT/'Integrated_Review_Binay_Verma.docx'),
        manuscript_sha256=sha(OUT/'MANUSCRIPT.md'), audit_source_sha256=sha(__file__))
    (OUT/'independent_reporting_audit.json').write_text(json.dumps(audit, indent=2)+'\n')
    print(json.dumps(audit, indent=2))


if __name__ == '__main__':
    main()
